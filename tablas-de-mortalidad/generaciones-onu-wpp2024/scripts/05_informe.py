"""Extrae las cifras del libro recalculado y genera el informe Word (llama a scripts/05_informe.js).

Uso:  python3 scripts/05_informe.py
"""
import json
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
wb = load_workbook(RAIZ / "resultados" / "tablas_mortalidad_generaciones_mexico.xlsx", data_only=True)
v, c = wb["Validacion"], wb["Comparacion"]

GENS = [1950, 1960, 1970, 1980, 1990, 2000]
SEXOS = ["Hombres", "Mujeres", "Ambos sexos"]
datos = {"gens": GENS, "sexos": SEXOS, "periodo": [], "gen": {}, "comp": {}}

for j, s in enumerate(SEXOS):
    datos["periodo"].append({"sexo": s, "media": v.cell(6 + j, 2).value, "mae": v.cell(6 + j, 3).value,
                             "max": v.cell(6 + j, 4).value, "anio_max": v.cell(6 + j, 5).value,
                             "d1950": v.cell(6 + j, 6).value, "d2023": v.cell(6 + j, 7).value,
                             "d2100": v.cell(6 + j, 8).value})
    datos["gen"][s] = []
    for i, g in enumerate(GENS):
        r = 13 + 6 * j + i
        datos["gen"][s].append({"g": g, "libro": v.cell(r, 3).value, "onu": v.cell(r, 4).value,
                                "dif": v.cell(r, 5).value, "periodo": v.cell(r, 7).value,
                                "brecha": v.cell(r, 8).value, "l100": v.cell(r, 9).value,
                                "ajuste": v.cell(r, 10).value,
                                "e0": c.cell(7 + i, 2 + 4 * j).value, "e65": c.cell(7 + i, 3 + 4 * j).value,
                                "l65": c.cell(7 + i, 4 + 4 * j).value, "l80": c.cell(7 + i, 5 + 4 * j).value})

salida = RAIZ / "resultados" / "informe_tablas_mortalidad_generaciones.docx"
with tempfile.TemporaryDirectory() as tmp:
    ruta = Path(tmp) / "datos_informe.json"
    ruta.write_text(json.dumps(datos, ensure_ascii=False))
    subprocess.run(["node", str(RAIZ / "scripts" / "05_informe.js"), str(ruta), str(RAIZ / "resultados"), str(salida)],
                   check=True)
print("informe", salida)
