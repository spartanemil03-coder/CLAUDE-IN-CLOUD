"""Genera los gráficos PNG a partir de los valores calculados del libro de Excel (ya recalculado).

Uso:  python3 scripts/03_graficos.py
"""
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
LIBRO = RAIZ / "resultados" / "tablas_mortalidad_generaciones_mexico.xlsx"
GENS = [1950, 1960, 1970, 1980, 1990, 2000]
SEXOS = [("H", "Hombres"), ("M", "Mujeres"), ("A", "Ambos sexos")]
BASE = {"H": 3, "M": 11, "A": 19}
OFF = {"lx": 3, "ex": 7}
RAMPA = ["#8fb8ea", "#6a9ede", "#3f82d0", "#2a64b0", "#1d4a87", "#12315c"]
C_SEXO = {"H": "#2a78d6", "M": "#eb6834", "A": "#1baf7a"}
INK, INK2, GRID = "#0b0b0b", "#52514e", "#e4e3df"

plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "axes.edgecolor": GRID, "axes.labelcolor": INK2,
                     "xtick.color": INK2, "ytick.color": INK2, "text.color": INK, "axes.spines.top": False,
                     "axes.spines.right": False, "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.8,
                     "axes.axisbelow": True, "figure.facecolor": "#fcfcfb", "axes.facecolor": "#fcfcfb",
                     "savefig.facecolor": "#fcfcfb"})

wb = load_workbook(LIBRO, data_only=True)


def serie_gen(g, s, k):
    ws = wb[f"Gen_{g}"]
    c = BASE[s] + OFF[k]
    return np.array([ws.cell(6 + i, c).value for i in range(101)], dtype=float)


def guardar(fig, nombre):
    fig.savefig(RAIZ / "resultados" / nombre, dpi=160, bbox_inches="tight")
    plt.close(fig)
    print("png", nombre)


# 1 y 2: ex y lx por generación
for k, nombre, ylab, titulo in (("ex", "ex_por_generacion.png", "ex: años que se esperan vivir",
                                 "Esperanza de vida por edad (ex), por generación"),
                                ("lx", "lx_por_generacion.png", "lx: sobrevivientes de 100,000 nacidos",
                                 "Sobrevivientes por edad (lx), por generación")):
    fig, axs = plt.subplots(1, 3, figsize=(15, 4.8), sharey=True)
    for ax, (s, ns) in zip(axs, SEXOS):
        for g, c in zip(GENS, RAMPA):
            y = serie_gen(g, s, k)
            ax.plot(range(101), y, color=c, lw=2, label=f"Generación {g}")
        ax.set_title(ns, loc="left", fontsize=11, fontweight="bold")
        ax.set_xlabel("Edad x")
        ax.set_xlim(0, 100)
        ax.set_xticks(range(0, 101, 20))
        if k == "ex":
            ax.set_ylim(0, 85)
        else:
            ax.set_ylim(0, 100000)
            ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: f"{int(v):,}"))
    axs[0].set_ylabel(ylab)
    h, l = axs[0].get_legend_handles_labels()
    fig.legend(h, l, loc="lower center", ncol=6, frameon=False, bbox_to_anchor=(0.5, -0.04))
    fig.suptitle(titulo + " - México (ONU WPP 2024, variante media)", x=0.01, ha="left", fontsize=13, fontweight="bold")
    fig.tight_layout(rect=(0, 0.02, 1, 0.95))
    guardar(fig, nombre)

# 3: e0 por año calendario (ONU) + e0 de cada generación
ws = wb["ONU_e0"]
anios = np.array([ws.cell(5 + i, 1).value for i in range(151)])
fig, ax = plt.subplots(figsize=(11, 5.4))
for j, (s, ns) in enumerate(SEXOS):
    y = np.array([ws.cell(5 + i, 2 + j).value for i in range(151)], dtype=float)
    ax.plot(anios, y, color=C_SEXO[s], lw=2, label=f"{ns}: e0 oficial ONU (periodo)")
    yg = [serie_gen(g, s, "ex")[0] for g in GENS]
    ax.scatter(GENS, yg, s=70, color=C_SEXO[s], edgecolor="#fcfcfb", linewidth=1.5, zorder=3)
    ax.text(2101, y[-1], ns, color=INK, va="center", fontsize=10)
ax.axvline(2023.5, color=INK2, lw=1, ls=(0, (4, 3)))
ax.text(2024.5, 42, "proyección →", color=INK2, fontsize=9)
ax.scatter([], [], s=70, color=INK2, label="e0 de cada generación (tablas del libro), en su año de nacimiento")
ax.set_xlim(1950, 2115)
ax.set_ylim(40, 90)
ax.set_xlabel("Año calendario / año de nacimiento")
ax.set_ylabel("Esperanza de vida al nacer (años)")
ax.set_title("Esperanza de vida al nacer por año calendario: periodo (ONU) y generaciones - México",
             loc="left", fontsize=12, fontweight="bold")
ax.legend(loc="lower right", frameon=False, fontsize=9)
fig.tight_layout()
guardar(fig, "e0_por_anio_calendario.png")

# 4: validación de periodo (diferencia calculada - ONU)
wx = wb["Validacion"]
fig, axs = plt.subplots(1, 2, figsize=(14, 4.8), gridspec_kw={"width_ratios": [1.5, 1]})
ax = axs[0]
for j, (s, ns) in enumerate(SEXOS):
    cd = 4 + 3 * j
    d = np.array([wx.cell(42 + i, cd).value for i in range(151)], dtype=float)
    ax.plot(anios, d, color=C_SEXO[s], lw=1.8, label=f"{ns} (error abs. medio {wx.cell(6 + j, 3).value:.2f})")
ax.axhline(0, color=INK2, lw=1)
ax.axvline(2023.5, color=INK2, lw=1, ls=(0, (4, 3)))
ax.set_xlim(1950, 2100)
ax.set_xlabel("Año calendario")
ax.set_ylabel("e0 calculada − e0 ONU (años)")
ax.set_title("A. e0 de periodo: método del libro vs. oficial ONU", loc="left", fontsize=11, fontweight="bold")
ax.legend(frameon=False, fontsize=9, loc="lower right")

ax = axs[1]
w = 0.26
x = np.arange(len(GENS))
for j, (s, ns) in enumerate(SEXOS):
    d = [wx.cell(13 + 6 * j + i, 5).value for i in range(6)]
    ax.bar(x + (j - 1) * w, d, w - 0.03, color=C_SEXO[s], label=ns)
ax.axhline(0, color=INK2, lw=1)
ax.set_xticks(x)
ax.set_xticklabels(GENS)
ax.set_xlabel("Generación")
ax.set_ylabel("e0 tabla del libro − e0 diagonal ONU (años)")
ax.set_title("B. Generaciones: libro vs. diagonal ONU", loc="left", fontsize=11, fontweight="bold")
ax.legend(frameon=False, fontsize=9, loc="lower right")
fig.suptitle("Validación contra la ONU - México", x=0.01, ha="left", fontsize=13, fontweight="bold")
fig.tight_layout(rect=(0, 0, 1, 0.94))
guardar(fig, "validacion_e0.png")

# 5: e0 por generación: libro vs diagonal ONU vs e0 de periodo del año de nacimiento
fig, axs = plt.subplots(1, 3, figsize=(15, 4.8), sharey=True)
w = 0.26
for ax, (j, (s, ns)) in zip(axs, enumerate(SEXOS)):
    f0 = 13 + 6 * j
    libro = [wx.cell(f0 + i, 3).value for i in range(6)]
    diag = [wx.cell(f0 + i, 4).value for i in range(6)]
    per = [wx.cell(f0 + i, 7).value for i in range(6)]
    for dx, v, c, lab in ((-1, libro, "#2a78d6", "Tabla de la generación (libro)"),
                          (0, diag, "#eb6834", "Generación con qx oficial ONU"),
                          (1, per, "#1baf7a", "e0 de periodo ONU, año de nacimiento")):
        ax.bar(x + dx * w, v, w - 0.03, color=c, label=lab)
    for xi, v in zip(x - w, libro):
        ax.text(xi, v + 0.8, f"{v:.1f}", ha="center", fontsize=8, color=INK)
    ax.set_xticks(x)
    ax.set_xticklabels(GENS)
    ax.set_ylim(30, 90)
    ax.set_title(ns, loc="left", fontsize=11, fontweight="bold")
    ax.set_xlabel("Generación")
axs[0].set_ylabel("e0 (años)")
h, l = axs[0].get_legend_handles_labels()
fig.legend(h, l, loc="lower center", ncol=3, frameon=False, bbox_to_anchor=(0.5, -0.05))
fig.suptitle("e0 por generación: tablas del libro, referencia oficial ONU y e0 de periodo", x=0.01, ha="left",
             fontsize=13, fontweight="bold")
fig.tight_layout(rect=(0, 0.02, 1, 0.95))
guardar(fig, "validacion_generaciones.png")
