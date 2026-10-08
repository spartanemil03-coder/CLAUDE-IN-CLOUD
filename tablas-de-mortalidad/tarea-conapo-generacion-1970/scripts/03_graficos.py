"""Gráficos PNG de la tabla de mortalidad de la generación 1970 (CONAPO), a partir de los valores calculados del libro.

Uso:  python3 scripts/03_graficos.py
"""
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.ticker import FuncFormatter
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import calculo_conapo as calc  # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent
LIBRO = RAIZ / "resultados" / "T_Mortalidad_CONAPO_Generacion_1970.xlsx"
G = calc.GENERACION
AZUL, NARANJA, VERDE = "#2a78d6", "#eb6834", "#1baf7a"
INK, INK2, GRID, BG = "#0b0b0b", "#52514e", "#e4e3df", "#fcfcfb"
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "axes.edgecolor": GRID, "axes.labelcolor": INK2,
                     "xtick.color": INK2, "ytick.color": INK2, "text.color": INK, "axes.spines.top": False,
                     "axes.spines.right": False, "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.8,
                     "axes.axisbelow": True, "figure.facecolor": BG, "axes.facecolor": BG, "savefig.facecolor": BG})
ws = load_workbook(LIBRO, data_only=True)["Tabla de mortalidad"]
T = {k: np.array([ws[f"{c}{5 + i}"].value for i in range(101)], float) for c, k in zip("BCDEFGH", ["x", "qx", "lx", "dx", "Lx", "Tx", "ex"])}
x = T["x"]
SUB = f"Generación mexicana de {G}, ambos sexos. Datos: CONAPO (conciliación 1950-2019 y proyecciones 2020-2070)"
ANIO_PROY = 2020 - G          # desde esta edad la generación entra en años proyectados (2020 en adelante)


def base(titulo, ylabel, xlim=(0, 100)):
    fig, ax = plt.subplots(figsize=(9, 5))
    ax.set_title(titulo, loc="left", fontsize=13, fontweight="bold", pad=22)
    ax.text(0, 1.02, SUB, transform=ax.transAxes, fontsize=8.5, color=INK2)
    ax.set_xlabel("Edad x (años cumplidos)")
    ax.set_ylabel(ylabel)
    ax.set_xlim(*xlim)
    return fig, ax


def zona_proyeccion(ax, xmax=100):
    ax.axvspan(ANIO_PROY, xmax, color="#eceae4", alpha=0.7, zorder=0)
    ax.axvline(ANIO_PROY, color=INK2, lw=0.9, ls=(0, (4, 3)))


def guardar(fig, nombre):
    fig.tight_layout()
    fig.savefig(RAIZ / "resultados" / nombre, dpi=160)
    plt.close(fig)
    print("png", nombre)


NOTA_PROY = f"≥ {ANIO_PROY} años: años calendario 2020-2070 (proyección)"

# 1. lx
fig, ax = base(f"Sobrevivientes lx de la generación de {G}", "lx (de 100,000 nacidos)")
zona_proyeccion(ax)
ax.plot(x, T["lx"], color=AZUL, lw=2.4)
ax.axhline(50000, color=INK2, lw=0.9, ls=(0, (2, 3)))
ax.text(1, 51500, "mitad de la generación", fontsize=8.5, color=INK2)
for edad in (65, 80):
    ax.scatter([edad], [T["lx"][edad]], color=AZUL, zorder=3, s=36)
    ax.annotate(f"{T['lx'][edad]:,.0f} llegan a los {edad}", (edad, T["lx"][edad]), xytext=(edad + 2, T["lx"][edad] + 7000), fontsize=9, color=INK)
ax.text(ANIO_PROY + 1, 97000, NOTA_PROY, fontsize=8.5, color=INK2)
ax.set_ylim(0, 102000)
ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{int(v):,}"))
guardar(fig, "lx_sobrevivientes.png")

# 2. qx (escala log)
fig, ax = base("Probabilidad de morir qx por edad (escala logarítmica)", "qx (log)", xlim=(0, 99))
zona_proyeccion(ax, 99)
ax.plot(x[:100], T["qx"][:100], color=AZUL, lw=2.4)
ax.set_yscale("log")
ax.set_ylim(4e-4, 0.5)
ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:g}"))
ax.grid(True, which="minor", color=GRID, lw=0.4)
ax.annotate(f"q0 = {T['qx'][0] * 1000:.1f} por mil", (0, T["qx"][0]), xytext=(6, T["qx"][0] * 0.9), fontsize=9, color=INK,
            arrowprops=dict(arrowstyle="-", color=INK2, lw=0.8))
ax.text(ANIO_PROY + 1, 6e-4, NOTA_PROY, fontsize=8.5, color=INK2)
ax.annotate("pandemia de COVID-19\n(años 2020-2021, edades 50-51)", (50.5, T["qx"][50]), xytext=(55, 0.03), fontsize=9, color=INK,
            arrowprops=dict(arrowstyle="-", color=INK2, lw=0.8))
guardar(fig, "qx_log.png")

# 3. dx
fig, ax = base("Defunciones dx de la generación por edad", "dx (de 100,000 nacidos)", xlim=(0, 99))
zona_proyeccion(ax, 99)
ax.plot(x[:100], T["dx"][:100], color=AZUL, lw=2.4)
ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{int(v):,}"))
ax.annotate(f"primer año: {T['dx'][0]:,.0f}", (0.5, T["dx"][0]), xytext=(7, T["dx"][0] * 0.95), fontsize=9, color=INK,
            arrowprops=dict(arrowstyle="-", color=INK2, lw=0.8))
moda = int(15 + np.argmax(T["dx"][15:100]))
ax.annotate(f"edad más frecuente de muerte adulta: {moda} años", (moda, T["dx"][moda]), xytext=(moda - 55, T["dx"][moda] + 250), fontsize=9, color=INK,
            arrowprops=dict(arrowstyle="-", color=INK2, lw=0.8))
ax.text(ANIO_PROY + 1, ax.get_ylim()[1] * 0.97, NOTA_PROY, fontsize=8.5, color=INK2)
ax.annotate("COVID-19 (2020-2021)", (51, T["dx"][51]), xytext=(54, 1500), fontsize=9, color=INK,
            arrowprops=dict(arrowstyle="-", color=INK2, lw=0.8))
guardar(fig, "dx_defunciones.png")

# 4. ex
fig, ax = base("Esperanza de vida ex por edad", "ex (años por vivir)")
zona_proyeccion(ax)
ax.plot(x, T["ex"], color=AZUL, lw=2.4)
for edad in (0, 65):
    ax.scatter([edad], [T["ex"][edad]], color=AZUL, zorder=3, s=36)
    ax.annotate(f"e{edad} = {T['ex'][edad]:.1f}", (edad, T["ex"][edad]), xytext=(edad + 3, T["ex"][edad] + 4), fontsize=9.5, color=INK)
ax.text(ANIO_PROY + 1, 80, NOTA_PROY, fontsize=8.5, color=INK2)
ax.set_ylim(0, 85)
guardar(fig, "ex_esperanza.png")

# 5. validación: e0 de periodo (mismas fórmulas) vs e0 oficial CONAPO, y e0 de la generación
D, P, ind = calc.cargar()
anios = np.arange(1970, 2071)
e_calc = np.array([calc.tabla_periodo_e0(D, P, ind.NAC[a], a) for a in anios])
e_conapo = ind.EV.reindex(anios).values
fig, ax = plt.subplots(figsize=(9, 5))
ax.set_title("e0 de periodo: método de la tabla vs. e0 oficial de CONAPO", loc="left", fontsize=13, fontweight="bold", pad=22)
ax.text(0, 1.02, "Mismas fórmulas aplicadas a cada año calendario (CONAPO: defunciones, población a mitad de año y nacimientos)", transform=ax.transAxes, fontsize=8.5, color=INK2)
ax.axvspan(2020, 2070, color="#eceae4", alpha=0.7, zorder=0)
ax.plot(anios, e_conapo, color=NARANJA, lw=3.2, label="e0 oficial CONAPO (EV)")
ax.plot(anios, e_calc, color=AZUL, lw=1.6, ls=(0, (4, 2)), label=f"e0 calculada (error abs. medio {np.abs(e_calc - e_conapo).mean():.2f} años)")
ax.scatter([G], [T["ex"][0]], color=VERDE, s=70, zorder=4, label=f"e0 de la generación {G} (esta tabla): {T['ex'][0]:.1f}")
ax.text(2023, 83.2, "proyección CONAPO", fontsize=8.5, color=INK2)
ax.annotate("COVID-19", (2020.5, 68.8), xytext=(2007, 66.5), fontsize=9, color=INK, arrowprops=dict(arrowstyle="-", color=INK2, lw=0.8))
ax.set_xlim(1970, 2070)
ax.set_xlabel("Año calendario")
ax.set_ylabel("e0 (años)")
ax.legend(frameon=False, loc="lower right")
guardar(fig, "validacion_e0_periodo.png")
