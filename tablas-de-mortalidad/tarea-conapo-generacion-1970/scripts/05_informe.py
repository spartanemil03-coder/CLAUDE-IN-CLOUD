"""Extrae las cifras del libro recalculado y genera el documento Word de análisis (llama a scripts/05_informe.js).

Uso:  python3 scripts/05_informe.py
"""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import calculo_conapo as calc  # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent
ws = load_workbook(RAIZ / "resultados" / "T_Mortalidad_CONAPO_Generacion_1970.xlsx", data_only=True)["Tabla de mortalidad"]
T = {k: np.array([ws[f"{c}{5 + i}"].value for i in range(101)], float) for c, k in zip("BCDEFGH", ["x", "qx", "lx", "dx", "Lx", "Tx", "ex"])}
D, P, ind = calc.cargar()
G = calc.GENERACION
dif = np.array([calc.tabla_periodo_e0(D, P, ind.NAC[a], a) - ind.EV[a] for a in range(1970, 2071)])
q, l, d, e = T["qx"], T["lx"], T["dx"], T["ex"]
datos = {
    "G": G,
    "e0": e[0], "e1": e[1], "e15": e[15], "e65": e[65], "e80": e[80],
    "q0_mil": q[0] * 1000, "5q0_mil": (1 - l[5] / l[0]) * 1000,
    "l15": l[15] / 1e5, "l65": l[65] / 1e5, "l80": l[80] / 1e5, "l100": l[100] / 1e5,
    "mediana": int(T["x"][np.where(l >= 5e4)[0][-1]]), "moda": int(15 + np.argmax(d[15:100])),
    "edad_supera_q0": int(1 + np.argmax(q[1:100] > q[0])), "q_min": float(q[1:30].min()), "edad_q_min": int(1 + np.argmin(q[1:30])),
    "q50": q[50], "q52": q[52], "q80": q[80], "q90": q[90], "q99": q[99],
    "nacimientos": float(ws["L6"].value), "def_inf": float(ws["L7"].value),
    "ev_1970": float(ind.EV[1970]), "ev_2019": float(ind.EV[2019]), "ev_2070": float(ind.EV[2070]),
    "mae": float(np.abs(dif).mean()), "maxdif": float(np.abs(dif).max()), "difmed": float(dif.mean()),
    "extracto": [[int(x), T["qx"][x], T["lx"][x], T["dx"][x], T["Lx"][x], T["Tx"][x], T["ex"][x]] for x in (0, 1, 5, 10, 20, 30, 40, 50, 60, 65, 70, 80, 90, 99, 100)],
}
salida = RAIZ / "resultados" / "Analisis_Tabla_Mortalidad_Generacion_1970.docx"
with tempfile.TemporaryDirectory() as tmp:
    ruta = Path(tmp) / "datos.json"
    ruta.write_text(json.dumps(datos, ensure_ascii=False))
    subprocess.run(["node", str(RAIZ / "scripts" / "05_informe.js"), str(ruta), str(RAIZ / "resultados"), str(salida)], check=True)
print("informe", salida)
