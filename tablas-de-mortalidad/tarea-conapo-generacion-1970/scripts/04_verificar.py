"""Verifica el libro recalculado contra el cálculo independiente (calculo_conapo.py) y los controles de las hojas de datos.

Uso:  python3 scripts/04_verificar.py        (AssertionError si algo no coincide)
"""
import sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import calculo_conapo as calc  # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent
wb = load_workbook(RAIZ / "resultados" / "T_Mortalidad_CONAPO_Generacion_1970.xlsx", data_only=True)
ws = wb["Tabla de mortalidad"]
D, P, ind = calc.cargar()
q = calc.qx_generacion(D, P, ind.NAC[calc.GENERACION])
t = calc.tabla(q)

# 1. columnas de la tabla
maxdif = 0.0
for col, k in zip("BCDEFGH", ["x", "qx", "lx", "dx", "Lx", "Tx", "ex"]):
    xl = np.array([ws[f"{col}{5 + i}"].value for i in range(101)], float)
    dif = float(np.max(np.abs(xl - t[k].values) / np.maximum(1, np.abs(t[k].values))))
    maxdif = max(maxdif, dif)
    assert dif < 1e-9, f"{k}: dif relativa {dif}"
# 2. controles de la plantilla
assert ws["L23"].value == 0, "suma de dx distinta del radix"
assert ws["L6"].value == ind.NAC[calc.GENERACION] and ws["L7"].value == D.loc[0, calc.GENERACION]
assert abs(ws["L9"].value - q[0]) < 1e-12
assert ws["D5"].value == 100000
# 3. hojas de datos: las matrices del libro son las de las bases; los totales cuadran con CONAPO
wn = wb["Nacimientos y e0"]
for i in range(101):
    r = 6 + i
    assert wn[f"E{r}"].value == 0 and wn[f"H{r}"].value == 0, f"totales no cuadran en {wn[f'A{r}'].value}"
wd = wb["Datos usados"]
for i in range(100):
    r = 5 + i
    assert wd[f"B{r}"].value == calc.GENERACION + i
    assert wd[f"C{r}"].value == D.loc[i, calc.GENERACION + i] and wd[f"D{r}"].value == P.loc[i, calc.GENERACION + i]
wdef = wb["Defunciones"]
assert all(wdef.cell(6 + e, 2 + j).value == D.loc[e, 1970 + j] for e in (0, 50, 109) for j in (0, 50, 100))
# 4. validación del método: e0 de periodo con las mismas fórmulas vs e0 oficial de CONAPO (columna EV)
difs = []
for a in range(1970, 2071):
    e0 = calc.tabla_periodo_e0(D, P, ind.NAC[a], a)
    difs.append(e0 - ind.EV[a])
difs = np.array(difs)
print(f"OK: libro verificado contra el cálculo independiente (dif. relativa máx {maxdif:.2e}); totales de datos cuadran con CONAPO")
print(f"Validación de periodo (1970-2070): dif. media {difs.mean():+.3f}, error absoluto medio {np.abs(difs).mean():.3f}, máx |dif| {np.abs(difs).max():.3f} años")
print(f"Generación {calc.GENERACION}: e0 = {t.ex[0]:.2f}, e65 = {t.ex[65]:.2f}, q0 = {q[0] * 1000:.2f} por mil, l65 = {t.lx[65] / 1e5:.3f}, l100 = {t.lx[100] / 1e5:.3f}")
