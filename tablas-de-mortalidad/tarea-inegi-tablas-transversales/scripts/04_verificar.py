"""Compara el libro de Excel (ya recalculado) con el cálculo independiente de calculo_abreviada.py.

Uso:  python3 scripts/04_verificar.py [libro.xlsx]
"""
import sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import calculo_abreviada as ca  # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent
LIBRO = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / "resultados" / "Tablas_transversales_abreviadas_INEGI_2025.xlsx"
wb = load_workbook(LIBRO, data_only=True)
COLS = {"P": "E", "D": "F", "nMx": "G", "nqx": "I", "lx": "K", "ndx": "L", "nLx": "M", "Tx": "N", "ex": "O"}
peor = 0.0
for s, hoja in (("H", "Hombres"), ("M", "Mujeres"), ("A", "Ambos sexos")):
    ref = ca.tabla(s)
    ws = wb[hoja]
    for campo, col in COLS.items():
        xl = np.array([ws[f"{col}{9 + i}"].value for i in range(17)], float)
        py = ref[campo].values.astype(float)
        err = np.nanmax(np.abs(xl - py) / np.maximum(np.abs(py), 1e-12))
        peor = max(peor, err)
        if err > 1e-9:
            print(f"DIFERENCIA en {hoja}.{campo}: error relativo {err:.3e}")
    print(f"{hoja}: e0 Excel = {ws['O9'].value:.4f}  Python = {ref.ex[0]:.4f}   control sum(ndx)-l0 = {ws['L27'].value}")
ins = wb["Insumos"]
print("Control reparto de defunciones (debe ser 0):", ins["H31"].value)
print("Error relativo máximo Excel vs Python:", f"{peor:.2e}")
print("OK" if peor < 1e-9 else "REVISAR")
