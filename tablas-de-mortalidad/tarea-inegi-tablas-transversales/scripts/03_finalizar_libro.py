"""Termina el libro de la profesora a partir de resultados/_sin_terminar.xlsx (salida de 02_construir_libro.py):

  1. Recalcula una COPIA con LibreOffice para obtener los valores de todas las fórmulas.
  2. Guarda esos valores dentro de las celdas de fórmula del libro (así se ven aunque el visor no recalcule); Excel igual recalcula al abrir.
  3. Restaura las capturas de pantalla originales de la hoja 'Fuentes de datos' (openpyxl pierde la imagen .emf).

El libro NO se pasa por LibreOffice, así que conserva el formato y las imágenes de la profesora.

Uso:  python3 scripts/03_finalizar_libro.py recalc.py [libro_sin_terminar.xlsx] [libro_final.xlsx]
      (recalc.py es el script recalc.py de la skill xlsx)
"""
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
RECALC = Path(sys.argv[1])
ENTRADA = Path(sys.argv[2]) if len(sys.argv) > 2 else RAIZ / "resultados" / "_sin_terminar.xlsx"
SALIDA = Path(sys.argv[3]) if len(sys.argv) > 3 else RAIZ / "resultados" / "Fuentes_de_consulta.xlsx"
ORIGINAL = RAIZ / "fuentes" / "Fuentes_de_consulta.xlsx"
HOJA_PROFE = "Fuentes de datos"

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PNS = "http://schemas.openxmlformats.org/package/2006/relationships"

# ---------- 1. valores calculados ----------
with tempfile.TemporaryDirectory() as tmp:
    copia = Path(tmp) / "copia.xlsx"
    shutil.copy(ENTRADA, copia)
    res = subprocess.run([sys.executable, str(RECALC), str(copia), "120"], capture_output=True, text=True)
    print(res.stdout.strip())
    if '"status": "success"' not in res.stdout:
        sys.exit("El recálculo con LibreOffice reportó errores; no se termina el libro.")
    valores = load_workbook(copia, data_only=True)
    VAL = {ws.title: {c.coordinate: c.value for fila in ws.iter_rows() for c in fila if c.value is not None} for ws in valores}

# ---------- 2 y 3. reescribir el paquete ----------
zin = zipfile.ZipFile(ENTRADA)
zorig = zipfile.ZipFile(ORIGINAL)
arch = {n: zin.read(n) for n in zin.namelist()}

wbx = etree.fromstring(arch["xl/workbook.xml"])
rels = etree.fromstring(arch["xl/_rels/workbook.xml.rels"])
destino = {r.get("Id"): r.get("Target") for r in rels}
hoja_xml = {}
for s in wbx.find(f"{{{NS}}}sheets"):
    t = destino[s.get(f"{{{RNS}}}id")]
    hoja_xml[s.get("name")] = t.lstrip("/") if t.startswith("/") else "xl/" + t

n_formulas = 0
for nombre, ruta in hoja_xml.items():
    if nombre == HOJA_PROFE:
        continue
    raiz = etree.fromstring(arch[ruta])
    for c in raiz.iter(f"{{{NS}}}c"):
        f = c.find(f"{{{NS}}}f")
        if f is None:
            continue
        v = c.find(f"{{{NS}}}v")
        if v is None:
            v = etree.SubElement(c, f"{{{NS}}}v")
        val = VAL[nombre].get(c.get("r"))
        if "t" in c.attrib:
            del c.attrib["t"]
        if isinstance(val, bool):
            c.set("t", "b")
            v.text = "1" if val else "0"
        elif isinstance(val, (int, float)):
            v.text = repr(float(val)) if isinstance(val, float) else str(val)
        else:
            c.set("t", "str")
            v.text = "" if val is None else str(val)
        n_formulas += 1
    arch[ruta] = etree.tostring(raiz, xml_declaration=True, encoding="UTF-8", standalone=True)
print("fórmulas con valor guardado:", n_formulas)

# capturas de la profesora: dibujo, sus relaciones y las imágenes originales
s1 = etree.fromstring(arch[f"xl/worksheets/_rels/{Path(hoja_xml[HOJA_PROFE]).name}.rels"])
dib = [r.get("Target") for r in s1 if r.get("Type").endswith("/drawing")][0]
dib = dib.lstrip("/") if dib.startswith("/") else str(Path("xl/worksheets") / dib).replace("xl/worksheets/../", "xl/")
nombre_dib = Path(dib).name
for n in [n for n in arch if n.startswith("xl/media/")]:
    del arch[n]
arch[f"xl/drawings/{nombre_dib}"] = zorig.read("xl/drawings/drawing1.xml")
arch[f"xl/drawings/_rels/{nombre_dib}.rels"] = zorig.read("xl/drawings/_rels/drawing1.xml.rels")
for n in zorig.namelist():
    if n.startswith("xl/media/"):
        arch[n] = zorig.read(n)
ct = arch["[Content_Types].xml"].decode()
if 'Extension="emf"' not in ct:
    ct = ct.replace("<Default ", '<Default Extension="emf" ContentType="image/x-emf"/><Default ', 1)
arch["[Content_Types].xml"] = ct.encode()

SALIDA.parent.mkdir(parents=True, exist_ok=True)
with zipfile.ZipFile(SALIDA, "w", zipfile.ZIP_DEFLATED) as zout:
    zout.writestr("[Content_Types].xml", arch.pop("[Content_Types].xml"))
    for n, b in arch.items():
        zout.writestr(n, b)
print("Libro final:", SALIDA)
if ENTRADA.name == "_sin_terminar.xlsx":
    ENTRADA.unlink()
