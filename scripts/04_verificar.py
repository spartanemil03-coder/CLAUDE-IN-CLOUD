"""Verifica el libro recalculado contra un cálculo independiente en numpy a partir de los CSV de datos/.

Uso:  python3 scripts/04_verificar.py        (falla con AssertionError si algo no coincide)
"""
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
wb = load_workbook(RAIZ / "resultados" / "tablas_mortalidad_generaciones_mexico.xlsx", data_only=True)
pop = pd.read_csv(RAIZ / "datos/mexico_PopulationBySingleAgeSex.csv")
dea = pd.read_csv(RAIZ / "datos/mexico_DeathsBySingleAgeSex.csv")
ind = pd.read_csv(RAIZ / "datos/mexico_DemographicIndicators.csv").set_index("Time")
tv = pd.read_csv(RAIZ / "datos/mexico_ONU_TablaVida_Completa.csv")
for d in (pop, dea):
    d["edad"] = d.AgeGrpStart.astype(int)

CP = {"H": "PopMale", "M": "PopFemale", "A": "PopTotal"}
CD = {"H": "DeathMale", "M": "DeathFemale", "A": "DeathTotal"}
SX = {"H": "Male", "M": "Female", "A": "Total"}
LEX = {"H": "LExMale", "M": "LExFemale", "A": "LEx"}
BASE = {"H": 3, "M": 11, "A": 19}
GENS = [1950, 1960, 1970, 1980, 1990, 2000]
RADIX = 100_000.0
maxdif = 0.0


def tabla(q):
    l = np.zeros(101)
    l[0] = RADIX
    for x in range(100):
        l[x + 1] = l[x] * (1 - q[x])
    d = l * q
    L = np.append(l[1:], 0.0) + d / 2
    T = L[::-1].cumsum()[::-1]
    return dict(qx=q, lx=l, dx=d, Lx=L, Tx=T, ex=T / l)


def cmp(a, b, etiqueta, tol=1e-6):
    global maxdif
    dif = float(np.max(np.abs(np.asarray(a, float) - np.asarray(b, float)) / np.maximum(1, np.abs(b))))
    maxdif = max(maxdif, dif)
    assert dif < tol, f"{etiqueta}: dif relativa máx {dif}"


for s in "HMA":
    P = pop.pivot(index="edad", columns="Time", values=CP[s])
    D = dea.pivot(index="edad", columns="Time", values=CD[s])
    Q = D / (P + D / 2)
    Q.loc[100] = 1.0
    for g in GENS:
        ws = wb[f"Gen_{g}"]
        t = tabla(np.array([Q.loc[x, g + x] for x in range(101)]))
        assert abs(t["lx"].sum() * 0 + t["dx"].sum() - RADIX) < 1e-6          # Σdx = radix
        for k, off in (("qx", 2), ("lx", 3), ("dx", 4), ("Lx", 5), ("Tx", 6), ("ex", 7)):
            xl = [ws.cell(6 + x, BASE[s] + off).value for x in range(101)]
            cmp(xl, t[k], f"Gen_{g} {s} {k}")
        # control de la hoja: Σdx - radix = 0
        assert abs(ws.cell(3, BASE[s] + 6).value) < 1e-6
    # periodo: e0 por año vs cálculo independiente y vs LEx oficial
    wp, wx = wb["Periodo"], wb["Validacion"]
    e0 = np.array([tabla(Q[a].values)["ex"][0] for a in range(1950, 2101)])
    j = "HMA".index(s)
    cmp([wp.cell(6 + j, 2 + i).value for i in range(151)], e0, f"Periodo e0 {s}")
    cmp([wx.cell(42 + i, 3 + 3 * j).value for i in range(151)], e0, f"Validacion calculada {s}")
    dif = e0 - ind.loc[1950:2100, LEX[s]].values
    cmp(wx.cell(6 + j, 3).value, np.abs(dif).mean(), f"MAE {s}")
    cmp(wx.cell(6 + j, 4).value, np.abs(dif).max(), f"max {s}")
    # diagonal ONU
    u = tv[tv.Sex == SX[s]].pivot(index="AgeGrpStart", columns="Time", values="qx")
    for i, g in enumerate(GENS):
        qd = np.array([u.loc[x, g + x] for x in range(101)])
        qd[100] = 1.0
        cmp(wx.cell(13 + 6 * j + i, 4).value, tabla(qd)["ex"][0], f"diagonal {s} {g}")
        cmp(wx.cell(13 + 6 * j + i, 3).value, tabla(np.array([Q.loc[x, g + x] for x in range(101)]))["ex"][0], f"e0 {s} {g}")

print(f"OK: libro verificado contra cálculo independiente (diferencia relativa máx {maxdif:.2e})")
