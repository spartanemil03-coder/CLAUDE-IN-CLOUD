"""Construye el libro de Excel a partir de la plantilla de la profesora (tarea_conapo/plantilla_profesora).

Hoja 1 'Tabla de mortalidad': la plantilla, completada de 0 a 100 años con fórmulas vivas (generación mexicana de 1970,
ambos sexos, datos de CONAPO). Hojas siguientes: los datos usados.

Uso:  python3 scripts/02_construir_excel.py    y luego recalcular con LibreOffice (recalc.py de la skill xlsx).
"""
import sys
from copy import copy
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.properties import PageSetupProperties

sys.path.insert(0, str(Path(__file__).resolve().parent))
import calculo_conapo as calc  # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent
PLANTILLA = RAIZ / "plantilla_profesora" / "T_Mortandad_1.xlsx"
SALIDA = RAIZ / "resultados" / "T_Mortalidad_CONAPO_Generacion_1970.xlsx"
G = calc.GENERACION
AZUL, VERDE, GRISTX = "0000FF", "008000", "595959"
ANIOS = list(range(1970, 2071))
EDADES = list(range(0, 110))

# ----------------------------------------------------------------------------- datos (agregados de las bases originales)
d_raw = pd.read_csv(calc.F_DEF)
p_raw = pd.read_csv(calc.F_POB)
ind = pd.read_csv(calc.F_IND)
ind = ind[ind.CVE_GEO == 0].set_index("ANIO")
D = d_raw.groupby(["EDAD", "ANIO"]).DEFUNCIONES.sum().unstack()
P = p_raw.groupby(["EDAD", "ANIO"]).POBLACION.sum().unstack()
assert D.shape == (110, 101) and P.shape == (110, 101)

wb = load_workbook(PLANTILLA)
ws = wb["Tabla de mortalidad"]


def f(**k):
    return Font(name="Calibri", size=k.pop("size", 10), **k)


def encabezado(c, texto):
    c.value = texto
    c.font = f(bold=True)
    c.fill = PatternFill("solid", fgColor="F2F2F2")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def nota(hoja, ref, texto, size=9):
    hoja[ref] = texto
    hoja[ref].font = f(italic=True, size=size, color=GRISTX)


# ============================================================================ Hoja 1: plantilla completada
for rng in list(ws.merged_cells.ranges):
    if str(rng) in ("B26:H26", "B33:H33"):
        ws.unmerge_cells(str(rng))
COLS = "BCDEFGH"
estilo = {n: {c: copy(ws[f"{c}{n}"]._style) for c in COLS} for n in (5, 6, 32)}
for r in range(5, 40):                       # limpiar el cuerpo de la plantilla (valores y comentarios fijos)
    for c in COLS:
        ws[f"{c}{r}"].value = None
        ws[f"{c}{r}"].comment = None
        ws[f"{c}{r}"]._style = copy(ws["A60"]._style)
ws["B2"] = f"Cuadro: Tabla de mortalidad de la generación mexicana de {G}, ambos sexos (CONAPO)"

P0, P1 = 5, 105
for i in range(101):
    r = P0 + i
    fila = 5 if r == P0 else 32 if r == P1 else 6
    for c in COLS:
        ws[f"{c}{r}"]._style = copy(estilo[fila][c])
    ws[f"B{r}"] = i
    if i == 0:
        ws[f"C{r}"] = "=L9"
    elif i == 100:
        ws[f"C{r}"] = 1
    else:
        ws[f"C{r}"] = f"='Datos usados'!C{5 + i}/('Datos usados'!D{5 + i}+'Datos usados'!C{5 + i}/2)"
    ws[f"D{r}"] = 100000 if i == 0 else f"=D{r-1}-E{r-1}"
    ws[f"E{r}"] = f"=ROUND(D{r}*C{r},0)"
    ws[f"F{r}"] = f"=ROUND(E{r}/2,0)" if i == 100 else f"=ROUND(D{r+1}+(E{r}/2),0)"
    ws[f"G{r}"] = f"=SUM(F{r}:F${P1})"
    ws[f"H{r}"] = f"=G{r}/D{r}"
    ws[f"C{r}"].number_format = "0.000000"
    for c in "DEFG":
        ws[f"{c}{r}"].number_format = "#,##0"
    ws[f"B{r}"].number_format = '0"+"' if i == 100 else "General"
for c, w in zip("ABCDEFGH", (6.1, 5.4, 35.3, 26.6, 12.6, 22, 36, 20)):
    ws.column_dimensions[c].width = w
ws.column_dimensions["K"].width = 34
ws.column_dimensions["L"].width = 24
ws.freeze_panes = "C5"

# nota técnica (como en la plantilla; ahora con fórmulas) y datos de la generación
ws["L6"] = "=INDEX('Nacimientos y e0'!B6:B106,MATCH(L16,'Nacimientos y e0'!A6:A106,0))"
ws["L7"] = "='Datos usados'!C5"
ws["K7"] = "Defunciones (menores de 1 año)"
ws["L9"] = "=L7/L6"
for ref in ("L6", "L7"):
    ws[ref].number_format = "#,##0"
    ws[ref].font = Font(name="Calibri", size=11, color=VERDE)
ws["L9"].number_format = "0.000000"
ws["K10"] = "q(0) = defunciones de menores de 1 año / nacimientos"
ws["K11"] = "q(x), x = 1…99 = defunciones / (población a mitad de año + defunciones/2)"
ws["K16"] = "Generación (año de nacimiento)"
ws["L16"] = G
ws["L16"].font = Font(name="Calibri", size=11, bold=True, color=AZUL)
ws["L16"].fill = PatternFill("solid", fgColor="FFFF00")
ws["K16"].font = Font(name="Calibri", size=11)
ws["K17"] = "Año calendario de la edad x = generación + x"
ws["K19"] = "Tx = Lx + L(x+1) + … + L100"
ws["K20"] = "ex = Tx / lx"
ws["K21"] = "Edad 100 = grupo abierto (100 y más): qx = 1"
ws["K23"] = "Control: suma de dx − radix"
ws["L23"] = f"=SUM(E{P0}:E{P1})-D{P0}"
ws["L23"].number_format = "0;-0;0"
for ref in ("K10", "K11", "K17", "K19", "K20", "K21", "K23"):
    ws[ref].font = Font(name="Calibri", size=9, italic=True, color=GRISTX)
ws["K25"] = "Fuente: CONAPO, Conciliación demográfica de México 1950-2019 y Proyecciones de la población 2020-2070."
ws["K26"] = "Datos y bases usadas: hojas siguientes y carpeta bases_de_datos_CONAPO/."
for ref in ("K25", "K26"):
    ws[ref].font = Font(name="Calibri", size=9, italic=True, color=GRISTX)

# ============================================================================ Datos usados (diagonal de la generación)
wd = wb.create_sheet("Datos usados")
wd.sheet_view.showGridLines = False
wd["A1"] = f"Datos usados en la tabla: generación mexicana de {G}, ambos sexos"
wd["A1"].font = f(size=13, bold=True)
nota(wd, "A2", "Para cada edad x se toman las defunciones y la población a mitad de año del año g + x (hojas Defunciones y Población). "
               "Verde = vínculo a otra hoja; negro = fórmula.")
for j, t in enumerate(["Edad x", "Año calendario (g + x)", "Defunciones de la edad x", "Población a mitad de año de la edad x", "q(x) = D / (P + D/2)"]):
    encabezado(wd.cell(4, 1 + j), t)
wd.row_dimensions[4].height = 44
for i in range(101):
    r = 5 + i
    wd.cell(r, 1, i).font = f(bold=True)
    wd.cell(r, 1).alignment = Alignment(horizontal="center")
    wd.cell(r, 2, f"='Tabla de mortalidad'!$L$16+A{r}").font = f()
    if i < 100:
        col = f"MATCH($B{r},Defunciones!$B$5:$CX$5,0)"
        wd.cell(r, 3, f"=INDEX(Defunciones!$B$6:$CX$115,$A{r}+1,{col})")
        wd.cell(r, 4, f"=INDEX('Población'!$B$6:$CX$115,$A{r}+1,{col.replace('Defunciones', chr(39) + 'Población' + chr(39))})")
        wd.cell(r, 5, f"=C{r}/(D{r}+C{r}/2)" if i else "='Tabla de mortalidad'!C5")
        for c in (3, 4):
            wd.cell(r, c).font = f(color=VERDE)
            wd.cell(r, c).number_format = "#,##0"
        wd.cell(r, 5).number_format = "0.000000"
        wd.cell(r, 5).font = f()
    else:
        wd.cell(r, 2).number_format = "0"
        wd.cell(r, 3, "grupo abierto: qx = 1").font = f(italic=True, color=GRISTX)
wd["A108"] = "q(0) usa nacimientos (hoja Nacimientos y e0): q(0) = defunciones de menores de 1 año / nacimientos."
wd["A108"].font = f(italic=True, size=9, color=GRISTX)
wd.column_dimensions["A"].width = 9
for c, w in zip("BCDE", (20, 22, 26, 20)):
    wd.column_dimensions[c].width = w
wd.freeze_panes = "A5"


# ============================================================================ Matrices Defunciones y Población
def matriz(nombre, M, titulo, fuente):
    w = wb.create_sheet(nombre)
    w.sheet_view.showGridLines = False
    w["A1"] = titulo
    w["A1"].font = f(size=13, bold=True)
    nota(w, "A2", fuente)
    nota(w, "A3", "Cada celda = suma de los 32 estados y de hombres + mujeres. Azul = dato de CONAPO. La tabla usa la diagonal edad x, año g + x.")
    encabezado(w["A5"], "Edad \\ Año")
    for j, a in enumerate(ANIOS):
        c = w.cell(5, 2 + j, a)
        encabezado(c, a)
    for i, e in enumerate(EDADES):
        w.cell(6 + i, 1, e).font = f(bold=True)
        w.cell(6 + i, 1).alignment = Alignment(horizontal="center")
        for j, a in enumerate(ANIOS):
            c = w.cell(6 + i, 2 + j, int(M.loc[e, a]))
            c.font = f(color=AZUL)
            c.number_format = "#,##0"
    w.freeze_panes = "B6"
    w.column_dimensions["A"].width = 11
    for j in range(len(ANIOS)):
        w.column_dimensions[L(2 + j)].width = 11
    return w


matriz("Defunciones", D, "Defunciones por edad y año, México (1970-2070)",
       "Fuente: CONAPO, 01_Defunciones_1950_2070.csv (carpeta bases_de_datos_CONAPO/).")
matriz("Población", P, "Población a mitad de año por edad y año, México (1970-2070)",
       "Fuente: CONAPO, 00_Pob_Mitad_1950_2070.csv (carpeta bases_de_datos_CONAPO/).")

# ============================================================================ Nacimientos y e0
wn = wb.create_sheet("Nacimientos y e0")
wn.sheet_view.showGridLines = False
wn["A1"] = "Nacimientos y otros indicadores nacionales de CONAPO, 1970-2070"
wn["A1"].font = f(size=13, bold=True)
nota(wn, "A2", "Fuente: CONAPO, 05_indicadores_demograficos_proyecciones.csv, fila 'República Mexicana'. Las columnas de verificación "
               "comparan los totales de CONAPO con la suma de las matrices de Defunciones y Población.")
heads = ["Año", "Nacimientos (NAC)", "Defunciones totales (DEF)", "Defunciones: suma de la hoja Defunciones", "Diferencia",
         "Población a mitad de año (POB_MIT_ANIO)", "Población: suma de la hoja Población", "Diferencia",
         "Mortalidad infantil (TMI, por mil)", "e0 total (EV)", "e0 hombres (EVH)", "e0 mujeres (EVM)"]
for j, t in enumerate(heads):
    encabezado(wn.cell(5, 1 + j), t)
wn.row_dimensions[5].height = 58
for i, a in enumerate(ANIOS):
    r = 6 + i
    col = L(2 + i)
    wn.cell(r, 1, a).font = f(bold=True)
    valores = {2: ind.NAC[a], 3: ind.DEF[a], 6: ind.POB_MIT_ANIO[a], 9: ind.TMI[a], 10: ind.EV[a], 11: ind.EVH[a], 12: ind.EVM[a]}
    for c, v in valores.items():
        wn.cell(r, c, float(v) if c >= 9 else int(v)).font = f(color=AZUL)
    wn.cell(r, 4, f"=SUM(Defunciones!{col}6:{col}115)").font = f()
    wn.cell(r, 5, f"=D{r}-C{r}").font = f()
    wn.cell(r, 7, f"=SUM('Población'!{col}6:{col}115)").font = f()
    wn.cell(r, 8, f"=G{r}-F{r}").font = f()
    for c in (2, 3, 4, 5, 6, 7, 8):
        wn.cell(r, c).number_format = "#,##0"
    for c in (9, 10, 11, 12):
        wn.cell(r, c).number_format = "0.00"
wn.freeze_panes = "B6"
wn.column_dimensions["A"].width = 8
for j in range(2, 13):
    wn.column_dimensions[L(j)].width = 17

# ============================================================================ Fuentes
wf = wb.create_sheet("Fuentes")
wf.sheet_view.showGridLines = False
wf["A1"] = "Bases de datos usadas (carpeta bases_de_datos_CONAPO/)"
wf["A1"].font = f(size=13, bold=True)
nota(wf, "A2", "Fuente: Consejo Nacional de Población (CONAPO), Conciliación demográfica de México 1950-2019 y Proyecciones de la población de México "
               "y de las entidades federativas 2020-2070. Licencia CC BY 4.0. Conjunto de datos: https://www.datos.gob.mx/dataset/proyecciones-de-poblacion")
cols = ["Archivo (tal como se descargó)", "Qué contiene", "Columnas usadas", "Para qué se usó", "Hoja donde aparece"]
for j, t in enumerate(cols):
    encabezado(wf.cell(4, 1 + j), t)
filas = [
    ("01_Defunciones_1950_2070.csv", "Defunciones por año (1970-2070), entidad, sexo y edad simple (0-109). Solo 32 entidades.",
     "ANIO, CVE_GEO, SEXO, EDAD, DEFUNCIONES", "Defunciones de cada edad en el año g + x (suma de estados y sexos).", "Defunciones; Datos usados"),
    ("00_Pob_Mitad_1950_2070.csv", "Población a mitad de año (1 de julio) por año (1970-2070), entidad, sexo y edad simple (0-109).",
     "ANIO, CVE_GEO, SEXO, EDAD, POBLACION", "Población expuesta al riesgo de cada edad (denominador de qx).", "Población; Datos usados"),
    ("05_indicadores_demograficos_proyecciones.csv", "Indicadores nacionales y estatales por año (1950-2070).",
     "ANIO, CVE_GEO = 0, NAC, DEF, POB_MIT_ANIO, TMI, EV, EVH, EVM", "Nacimientos del año de la generación (q0), verificación de totales y e0 oficial.", "Nacimientos y e0; Tabla de mortalidad"),
]
for i, fila in enumerate(filas):
    for j, v in enumerate(fila):
        c = wf.cell(5 + i, 1 + j, v)
        c.font = f(bold=(j == 0))
        c.alignment = Alignment(wrap_text=True, vertical="top")
    wf.row_dimensions[5 + i].height = 48
wf["A9"] = "Cómo se agregó"
wf["A9"].font = f(bold=True)
wf["A10"] = ("CONAPO publica defunciones y población por entidad, sexo y edad. Los datos nacionales de este libro son la suma de las 32 entidades y de "
             "hombres y mujeres. La suma coincide con los totales nacionales del archivo 05 (columna de verificación en la hoja Nacimientos y e0).")
wf["A10"].alignment = Alignment(wrap_text=True, vertical="top")
wf.merge_cells("A10:E10")
wf.row_dimensions[10].height = 44
wf["A10"].font = f()
for c, w in zip("ABCDE", (40, 52, 36, 46, 32)):
    wf.column_dimensions[c].width = w

for w in wb.worksheets:                     # impresión: apaisado y ajustado al ancho de una página
    w.page_setup.orientation = "landscape"
    w.sheet_properties.pageSetUpPr = w.sheet_properties.pageSetUpPr or PageSetupProperties()
    w.sheet_properties.pageSetUpPr.fitToPage = True
    w.page_setup.fitToWidth = 1
    w.page_setup.fitToHeight = 0
wb["Tabla de mortalidad"].page_setup.orientation = "portrait"
wb["Tabla de mortalidad"].print_title_rows = "3:4"

SALIDA.parent.mkdir(exist_ok=True)
wb.save(SALIDA)
print("guardado", SALIDA)
