"""Construye el libro de Excel a partir de la plantilla de la profesora (plantilla_profesora/).

Hoja 1 'Tabla de mortalidad': la plantilla, completada de 0 a 100 años con fórmulas vivas (México 2019, ambos sexos,
datos de INEGI). Hojas siguientes: los datos usados (defunciones, población, nacimientos) y las fuentes.

Uso:  python3 scripts/02_construir_excel.py    y luego recalcular con LibreOffice (recalc.py de la skill xlsx).
"""
import sys
from copy import copy
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties

sys.path.insert(0, str(Path(__file__).resolve().parent))
import calculo_inegi as calc  # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent
PLANTILLA = RAIZ / "plantilla_profesora" / "T_Mortandad_1.xlsx"
SALIDA = RAIZ / "resultados" / "T_Mortalidad_INEGI_Mexico_2019.xlsx"
AZUL, VERDE, GRISTX = "0000FF", "008000", "595959"

dfn, nac, pob = calc.cargar()
d_arch = pd.read_csv(calc.DATOS / "inegi_defunciones_2019_por_archivo.csv")
n_arch = pd.read_csv(calc.DATOS / "inegi_nacimientos_2019_por_archivo.csv")

wb = load_workbook(PLANTILLA)
ws = wb["Tabla de mortalidad"]


def f(**k):
    return Font(name="Calibri", size=k.pop("size", 10), **k)


def encabezado(c, texto):
    c.value = texto
    c.font = f(bold=True)
    c.fill = PatternFill("solid", fgColor="F2F2F2")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def nota(hoja, ref, texto, size=9):
    hoja[ref] = texto
    hoja[ref].font = f(italic=True, size=size, color=GRISTX)


# ============================================================================ Hoja 1: plantilla completada
for rng in list(ws.merged_cells.ranges):
    if str(rng) in ("B26:H26", "B33:H33"):
        ws.unmerge_cells(str(rng))
COLS = "BCDEFGH"
estilo = {n: {c: copy(ws[f"{c}{n}"]._style) for c in COLS} for n in (5, 6, 32)}
for r in range(5, 40):
    for c in COLS:
        ws[f"{c}{r}"].value = None
        ws[f"{c}{r}"].comment = None
        ws[f"{c}{r}"]._style = copy(ws["A60"]._style)
ws["B2"] = "Cuadro: Tabla de mortalidad de México, 2019, ambos sexos (INEGI)"

P0, P1 = 5, 105
for i in range(101):
    r = P0 + i
    fila = 5 if r == P0 else 32 if r == P1 else 6
    for c in COLS:
        ws[f"{c}{r}"]._style = copy(estilo[fila][c])
    ws[f"B{r}"] = i
    if i == 0:
        ws[f"C{r}"] = "=L9"
    elif i == 100:
        ws[f"C{r}"] = 1
    else:
        ws[f"C{r}"] = f"='Datos usados'!D{r}/('Datos usados'!F{r}+'Datos usados'!D{r}/2)"
    ws[f"D{r}"] = 100000 if i == 0 else f"=D{r-1}-E{r-1}"
    ws[f"E{r}"] = f"=ROUND(D{r}*C{r},0)"
    ws[f"F{r}"] = f"=ROUND(E{r}/2,0)" if i == 100 else f"=ROUND(D{r+1}+(E{r}/2),0)"
    ws[f"G{r}"] = f"=SUM(F{r}:F${P1})"
    ws[f"H{r}"] = f"=G{r}/D{r}"
    ws[f"C{r}"].number_format = "0.000000"
    for c in "DEFG":
        ws[f"{c}{r}"].number_format = "#,##0"
    ws[f"B{r}"].number_format = '0"+"' if i == 100 else "General"
for c, w in zip("ABCDEFGH", (6.1, 5.4, 35.3, 26.6, 12.6, 22, 36, 20)):
    ws.column_dimensions[c].width = w
ws.column_dimensions["K"].width = 38
ws.column_dimensions["L"].width = 24
ws.freeze_panes = "C5"

ws["L6"] = "=Nacimientos!F8"
ws["L7"] = "='Datos usados'!D5"
ws["K7"] = "Defunciones (menores de 1 año)"
ws["L9"] = "=L7/L6"
for ref in ("L6", "L7"):
    ws[ref].number_format = "#,##0"
    ws[ref].font = Font(name="Calibri", size=11, color=VERDE)
ws["L9"].number_format = "0.000000"
ws["K10"] = "q(0) = defunciones de menores de 1 año / nacimientos"
ws["K11"] = "q(x), x = 1…99 = defunciones / (población + defunciones/2)"
ws["K16"] = "Año de referencia"
ws["L16"] = 2019
ws["L16"].font = Font(name="Calibri", size=11, bold=True)
ws["K16"].font = Font(name="Calibri", size=11)
ws["K17"] = "Defunciones y nacimientos ocurridos en 2019; población del Censo 2020"
ws["K18"] = "Suavizado del Censo (1 = sí, 0 = no)"
ws["K18"].font = Font(name="Calibri", size=11)
ws["L18"] = 1
ws["L18"].font = Font(name="Calibri", size=11, bold=True, color=AZUL)
ws["L18"].fill = PatternFill("solid", fgColor="FFFF00")
ws["L18"].alignment = Alignment(horizontal="right")
ws["K19"] = "Media móvil de 5 edades en defunciones y población, de 10 a 97 años (hoja Datos usados)"
ws["K21"] = "Tx = Lx + L(x+1) + … + L100"
ws["K22"] = "ex = Tx / lx"
ws["K23"] = "Edad 100 = grupo abierto (100 y más): qx = 1"
ws["K25"] = "Control: suma de dx − radix"
ws["L25"] = f"=SUM(E{P0}:E{P1})-D{P0}"
ws["L25"].number_format = "0;-0;0"
for ref in ("K10", "K11", "K17", "K19", "K21", "K22", "K23", "K25"):
    ws[ref].font = Font(name="Calibri", size=9, italic=True, color=GRISTX)
ws["K27"] = "Fuente: INEGI (Estadística de Nacimientos Registrados, Estadísticas de Defunciones Registradas y Censo de Población y Vivienda 2020)."
ws["K28"] = "Datos y bases usadas: hojas siguientes y carpeta bases_de_datos_INEGI/."
for ref in ("K27", "K28"):
    ws[ref].font = Font(name="Calibri", size=9, italic=True, color=GRISTX)
SW = "'Tabla de mortalidad'!$L$18"

# ============================================================================ Defunciones
wdf = wb.create_sheet("Defunciones")
wdf.sheet_view.showGridLines = False
wdf["A1"] = "Defunciones ocurridas en 2019 por edad y sexo, México (INEGI)"
wdf["A1"].font = f(size=13, bold=True)
nota(wdf, "A2", "Fuente: INEGI, Estadísticas de Defunciones Registradas (EDR), datos abiertos 2019, 2020 y 2021 (carpeta bases_de_datos_INEGI/1_Defunciones_EDR). "
                "Se cuentan las defunciones con año de ocurrencia 2019 (anio_ocur = 2019) registradas en esos tres años.")
for j, t in enumerate(["Edad x", "Hombres", "Mujeres", "Sexo no especificado", "Total"]):
    encabezado(wdf.cell(4, 1 + j), t)
for i in range(101):
    r = 5 + i
    wdf.cell(r, 1, i).font = f(bold=True)
    wdf.cell(r, 1).alignment = Alignment(horizontal="center")
    if i == 100:
        wdf.cell(r, 1).number_format = '0"+"'
    for j, k in enumerate(["hombres", "mujeres", "sexo_no_especificado"]):
        c = wdf.cell(r, 2 + j, int(dfn.loc[str(i), k]))
        c.font, c.number_format = f(color=AZUL), "#,##0"
    wdf.cell(r, 5, f"=SUM(B{r}:D{r})").number_format = "#,##0"
wdf.cell(106, 1, "No especificada").font = f(bold=True)
for j, k in enumerate(["hombres", "mujeres", "sexo_no_especificado"]):
    c = wdf.cell(106, 2 + j, int(dfn.loc["NE", k]))
    c.font, c.number_format = f(color=AZUL), "#,##0"
wdf.cell(106, 5, "=SUM(B106:D106)").number_format = "#,##0"
wdf.cell(107, 1, "Total").font = f(bold=True)
for c in range(2, 6):
    col = chr(64 + c)
    wdf.cell(107, c, f"=SUM({col}5:{col}106)").font = f(bold=True)
    wdf.cell(107, c).number_format = "#,##0"
wdf.cell(108, 1, "Factor de redistribución").font = f(bold=True)
wdf["E108"] = "=E107/(E107-E106)"
wdf["E108"].number_format = "0.0000"
nota(wdf, "A110", "Factor = total / total con edad especificada: reparte las defunciones de edad no especificada entre todas las edades en proporción.")
for j, t in enumerate(["Archivo EDR", "Año de registro", "Defunciones ocurridas en 2019"]):
    encabezado(wdf.cell(4, 7 + j), t)
for i, fila in d_arch.iterrows():
    wdf.cell(5 + i, 7, fila.archivo).font = f(color=AZUL)
    wdf.cell(5 + i, 8, int(fila.anio_registro)).font = f(color=AZUL)
    c = wdf.cell(5 + i, 9, int(fila.defunciones_ocurridas_2019))
    c.font, c.number_format = f(color=AZUL), "#,##0"
wdf.cell(8, 7, "Total").font = f(bold=True)
wdf["I8"] = "=SUM(I5:I7)"
wdf["I8"].font, wdf["I8"].number_format = f(bold=True), "#,##0"
wdf["G9"] = "Control (debe ser 0): total por archivo − total por edad"
wdf["G9"].font = f(italic=True, size=9, color=GRISTX)
wdf["I9"] = "=I8-E107"
wdf["I9"].number_format = "0;-0;0"
wdf.freeze_panes = "A5"
wdf.column_dimensions["A"].width = 17
for c, w in zip("BCDEFGHI", (13, 13, 14, 13, 3, 58, 14, 18)):
    wdf.column_dimensions[c].width = w
wdf.row_dimensions[4].height = 32

# ============================================================================ Población
wp = wb.create_sheet("Población")
wp.sheet_view.showGridLines = False
wp["A1"] = "Población por edad y sexo, México, Censo 2020 (INEGI)"
wp["A1"].font = f(size=13, bold=True)
nota(wp, "A2", "Fuente: INEGI, Censo de Población y Vivienda 2020, tabulados del cuestionario básico, 'Población 3' (cpv2020_b_eum_01_poblacion.xlsx, carpeta "
               "bases_de_datos_INEGI/3_Censo_2020). Estados Unidos Mexicanos, edad desplegada.")
for j, t in enumerate(["Edad x", "Hombres", "Mujeres", "Total"]):
    encabezado(wp.cell(4, 1 + j), t)
for i in range(101):
    r = 5 + i
    wp.cell(r, 1, i).font = f(bold=True)
    wp.cell(r, 1).alignment = Alignment(horizontal="center")
    if i == 100:
        wp.cell(r, 1).number_format = '0"+"'
    for j, k in enumerate(["hombres", "mujeres"]):
        c = wp.cell(r, 2 + j, int(pob.loc[str(i), k]))
        c.font, c.number_format = f(color=AZUL), "#,##0"
    wp.cell(r, 4, f"=B{r}+C{r}").number_format = "#,##0"
wp.cell(106, 1, "No especificada").font = f(bold=True)
for j, k in enumerate(["hombres", "mujeres"]):
    c = wp.cell(106, 2 + j, int(pob.loc["NE", k]))
    c.font, c.number_format = f(color=AZUL), "#,##0"
wp.cell(106, 4, "=B106+C106").number_format = "#,##0"
wp.cell(107, 1, "Total").font = f(bold=True)
for c in range(2, 5):
    col = chr(64 + c)
    wp.cell(107, c, f"=SUM({col}5:{col}106)").font = f(bold=True)
    wp.cell(107, c).number_format = "#,##0"
nota(wp, "A109", "Población con edad no especificada: no se redistribuye (0.2 % del total). El Censo concentra población en edades redondas "
                 "(p. ej. 30 años: 2.37 millones; 31 años: 1.54 millones); por eso la hoja Datos usados permite suavizar.")
wp.freeze_panes = "A5"
wp.column_dimensions["A"].width = 17
for c in "BCD":
    wp.column_dimensions[c].width = 15

# ============================================================================ Nacimientos
wn = wb.create_sheet("Nacimientos")
wn.sheet_view.showGridLines = False
wn["A1"] = "Nacimientos ocurridos en 2019, México (INEGI)"
wn["A1"].font = f(size=13, bold=True)
nota(wn, "A2", "Fuente: INEGI, Estadística de Nacimientos Registrados (ENR), datos abiertos 2019, 2020 y 2021 (carpeta bases_de_datos_INEGI/2_Nacimientos_ENR). "
               "Se cuentan los nacimientos con año de nacimiento 2019 (ano_nac = 2019) registrados en esos tres años.")
for j, t in enumerate(["Archivo ENR", "Año de registro", "Hombres", "Mujeres", "Sexo no especificado", "Total"]):
    encabezado(wn.cell(4, 1 + j), t)
for i, fila in n_arch.iterrows():
    r = 5 + i
    wn.cell(r, 1, fila.archivo).font = f(color=AZUL)
    wn.cell(r, 2, int(fila.anio_registro)).font = f(color=AZUL)
    for j, k in enumerate(["hombres", "mujeres", "sexo_no_especificado"]):
        c = wn.cell(r, 3 + j, int(fila[k]))
        c.font, c.number_format = f(color=AZUL), "#,##0"
    wn.cell(r, 6, f"=SUM(C{r}:E{r})").number_format = "#,##0"
wn.cell(8, 1, "Total").font = f(bold=True)
for c in range(3, 7):
    col = chr(64 + c)
    wn.cell(8, c, f"=SUM({col}5:{col}7)").font = f(bold=True)
    wn.cell(8, c).number_format = "#,##0"
nota(wn, "A10", "Una parte importante de los nacimientos de 2019 se registró en 2020 y 2021 (registro tardío); por eso se usan los tres archivos.")
wn.column_dimensions["A"].width = 52
for c, w in zip("BCDEF", (16, 14, 14, 20, 14)):
    wn.column_dimensions[c].width = w
wn.row_dimensions[4].height = 32

# ============================================================================ Datos usados
wd = wb.create_sheet("Datos usados", 1)
wd.sheet_view.showGridLines = False
wd["A1"] = "Datos usados en la tabla de mortalidad, México 2019, ambos sexos"
wd["A1"].font = f(size=13, bold=True)
nota(wd, "A2", "Verde = vínculo a otra hoja; negro = fórmula. Las columnas 'usadas' aplican la media móvil de 5 edades (10 a 97 años) si el suavizado "
               "de la hoja Tabla de mortalidad está activado (1).")
heads = ["Edad x", "Defunciones ocurridas en 2019 (EDR)", "Defunciones con edad no especificada redistribuida", "Defunciones usadas en qx",
         "Población, Censo 2020", "Población usada en qx", "q(x) = D / (P + D/2)"]
for j, t in enumerate(heads):
    encabezado(wd.cell(4, 1 + j), t)
wd.row_dimensions[4].height = 58
for i in range(101):
    r = 5 + i
    wd.cell(r, 1, i).font = f(bold=True)
    wd.cell(r, 1).alignment = Alignment(horizontal="center")
    if i == 100:
        wd.cell(r, 1).number_format = '0"+"'
    wd.cell(r, 2, f"=Defunciones!E{r}").font = f(color=VERDE)
    wd.cell(r, 3, f"=B{r}*Defunciones!$E$108").font = f()
    wd.cell(r, 4, f"=IF(AND({SW}=1,$A{r}>=10,$A{r}<=97),AVERAGE(C{r-2}:C{r+2}),C{r})").font = f()
    wd.cell(r, 5, f"=Población!D{r}").font = f(color=VERDE)
    wd.cell(r, 6, f"=IF(AND({SW}=1,$A{r}>=10,$A{r}<=97),AVERAGE(E{r-2}:E{r+2}),E{r})").font = f()
    if i == 0:
        wd.cell(r, 7, "='Tabla de mortalidad'!C5").font = f(color=VERDE)
    elif i < 100:
        wd.cell(r, 7, f"=D{r}/(F{r}+D{r}/2)").font = f()
    else:
        wd.cell(r, 7, 1).font = f(color=AZUL)
    for c in range(2, 7):
        wd.cell(r, c).number_format = "#,##0.0" if c in (3, 4, 6) else "#,##0"
    wd.cell(r, 7).number_format = "0.000000"
nota(wd, "A107", "q(0) usa nacimientos (hoja Nacimientos): q(0) = defunciones de menores de 1 año / nacimientos. q(100) = 1 (grupo abierto).")
wd.freeze_panes = "A5"
wd.column_dimensions["A"].width = 9
for c, w in zip("BCDEFG", (22, 28, 22, 18, 20, 20)):
    wd.column_dimensions[c].width = w

# ============================================================================ Fuentes
wf = wb.create_sheet("Fuentes")
wf.sheet_view.showGridLines = False
wf["A1"] = "Bases de datos usadas (carpeta bases_de_datos_INEGI/)"
wf["A1"].font = f(size=13, bold=True)
nota(wf, "A2", "Todas son bases de INEGI, descargadas tal cual y guardadas sin modificar en la carpeta bases_de_datos_INEGI/.")
cols = ["Carpeta / archivo", "Qué es", "Campos usados", "Para qué se usó", "Enlace"]
for j, t in enumerate(cols):
    encabezado(wf.cell(4, 1 + j), t)
filas = [
    ("1_Defunciones_EDR/conjunto_de_datos_defunciones_registradas_2019_csv.zip (y 2020, 2021)",
     "Estadísticas de Defunciones Registradas (EDR). Datos abiertos: una fila por defunción registrada en el año.",
     "sexo, edad, anio_ocur", "Defunciones ocurridas en 2019 por edad y sexo (se usan los tres años de registro por registro tardío).",
     "https://www.inegi.org.mx/programas/edr/#datos_abiertos"),
    ("2_Nacimientos_ENR/conjunto_de_datos_natalidad_2019_csv.zip (y 2020, 2021)",
     "Estadística de Nacimientos Registrados (ENR). Datos abiertos: una fila por nacimiento registrado en el año.",
     "sexo, ano_nac", "Nacimientos ocurridos en 2019 (q0 = defunciones de menores de 1 año / nacimientos).",
     "https://www.inegi.org.mx/programas/natalidad/#datos_abiertos"),
    ("3_Censo_2020/cpv2020_b_eum_01_poblacion.xlsx",
     "Censo de Población y Vivienda 2020, tabulados del cuestionario básico (Población), hoja 03: población por edad desplegada y sexo.",
     "Estados Unidos Mexicanos: edad, hombres, mujeres", "Población por edad (denominador de qx para x = 1 a 99).",
     "https://www.inegi.org.mx/programas/ccpv/2020/#tabulados"),
]
for i, fila in enumerate(filas):
    for j, v in enumerate(fila):
        c = wf.cell(5 + i, 1 + j, v)
        c.font = f(bold=(j == 0))
        c.alignment = Alignment(wrap_text=True, vertical="top")
    wf.row_dimensions[5 + i].height = 62
wf["A9"] = "Cómo se condensaron"
wf["A9"].font = f(bold=True)
wf["A10"] = ("Los archivos de EDR y ENR traen un registro por defunción o nacimiento. Para este libro se contaron (script scripts/01_descargar_inegi.py): "
             "defunciones ocurridas en 2019 por edad y sexo (código de edad 1001-3999 = menores de 1 año; 4001-4120 = años cumplidos; 4998 = no especificada) "
             "y nacimientos ocurridos en 2019 por sexo. Los conteos están en las hojas Defunciones, Nacimientos y Población.")
wf["A10"].alignment = Alignment(wrap_text=True, vertical="top")
wf["A10"].font = f()
wf.merge_cells("A10:E10")
wf.row_dimensions[10].height = 64
for c, w in zip("ABCDE", (52, 50, 30, 46, 46)):
    wf.column_dimensions[c].width = w

for w in wb.worksheets:
    w.page_setup.orientation = "landscape"
    w.sheet_properties.pageSetUpPr = w.sheet_properties.pageSetUpPr or PageSetupProperties()
    w.sheet_properties.pageSetUpPr.fitToPage = True
    w.page_setup.fitToWidth = 1
    w.page_setup.fitToHeight = 0
wb["Tabla de mortalidad"].page_setup.orientation = "portrait"
wb["Tabla de mortalidad"].print_title_rows = "3:4"
orden = ["Tabla de mortalidad", "Datos usados", "Defunciones", "Población", "Nacimientos", "Fuentes"]
wb._sheets = [wb[n] for n in orden]

SALIDA.parent.mkdir(exist_ok=True)
wb.save(SALIDA)
print("guardado", SALIDA)
