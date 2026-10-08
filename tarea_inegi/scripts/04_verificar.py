"""Verifica el libro recalculado contra el cálculo independiente (calculo_inegi.py), los controles de las hojas de datos
y que el interruptor de suavizado funcione.

Uso:  python3 scripts/04_verificar.py        (AssertionError si algo no coincide)
"""
import subprocess
import sys
import tempfile
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import calculo_inegi as calc  # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent
LIBRO = RAIZ / "resultados" / "2_Tabla_de_mortalidad_INEGI_2019.xlsx"
PREP = RAIZ / "resultados" / "1_Preparacion_datos_INEGI_2019.xlsx"


def comparar(wb, suaviza):
    ws = wb["Tabla de mortalidad"]
    q, d, p, b = calc.qx("A", suaviza)
    t = calc.tabla(q)
    maxdif = 0.0
    for col, k in zip("BCDEFGH", ["x", "qx", "lx", "dx", "Lx", "Tx", "ex"]):
        xl = np.array([ws[f"{col}{5 + i}"].value for i in range(101)], float)
        dif = float(np.max(np.abs(xl - t[k].values) / np.maximum(1, np.abs(t[k].values))))
        maxdif = max(maxdif, dif)
        assert dif < 1e-9, f"{k}: dif relativa {dif}"
    assert ws["L25"].value == 0 and ws["D5"].value == 100000
    assert ws["L6"].value == b and abs(ws["L9"].value - q[0]) < 1e-12
    return maxdif, t


# libro 2: la tabla contra el cálculo independiente
wb = load_workbook(LIBRO, data_only=True)
maxdif, t = comparar(wb, True)
print(f"OK libro 2: tabla verificada contra el cálculo independiente (dif. relativa máx {maxdif:.2e}); control sum(dx) - radix = 0")
# libro 1: controles de las hojas de datos y de la preparación
wp = load_workbook(PREP, data_only=True)
assert wp["Preparación"]["J5"].value == 1, "el libro 1 debe entregarse con el suavizado activado"
assert wp["Defunciones"]["I9"].value == 0, "total de defunciones por archivo ≠ total por edad"
dfn, nac, pob = calc.cargar()
assert wp["Defunciones"]["E107"].value == dfn.values.sum()
assert wp["Nacimientos"]["F8"].value == nac.sum()
assert wp["Población"]["D107"].value == pob.values.sum() == 126_014_024
q, d, p, b = calc.qx("A", True)
for col, ref in (("D", d), ("F", p), ("G", q)):
    xl = np.array([wp["Preparación"][f"{col}{5 + i}"].value for i in range(101)], float)
    assert np.max(np.abs(xl - ref) / np.maximum(1e-12, np.abs(ref))) < 1e-9, col
# los insumos del libro 2 son iguales a las columnas D y F del libro 1
wi = wb["Insumos"]
for col, src in (("B", "D"), ("C", "F")):
    a = np.array([wi[f"{col}{5 + i}"].value for i in range(101)], float)
    b_ = np.array([wp["Preparación"][f"{src}{5 + i}"].value for i in range(101)], float)
    assert np.max(np.abs(a - b_)) < 1e-6, col
assert wi["F5"].value == wp["Nacimientos"]["F8"].value
print("OK libro 1: controles de datos en 0; sus columnas D y F coinciden con la hoja Insumos del libro 2")

# interruptor de suavizado (libro 1): se apaga en una copia, se recalcula con LibreOffice y se compara con la referencia sin suavizar
with tempfile.TemporaryDirectory() as tmp:
    copia = Path(tmp) / "prueba.xlsx"
    w = load_workbook(PREP)
    w["Preparación"]["J5"] = 0
    w.save(copia)
    skill = next(Path("/root/.claude/skills/synced").glob("*/xlsx/scripts/recalc.py"))
    subprocess.run([sys.executable, str(skill), str(copia), "300"], check=True, capture_output=True)
    w2 = load_workbook(copia, data_only=True)["Preparación"]
    q0, d0, p0, _ = calc.qx("A", False)
    for col, ref in (("D", d0), ("F", p0), ("G", q0)):
        xl = np.array([w2[f"{col}{5 + i}"].value for i in range(101)], float)
        assert np.max(np.abs(xl - ref) / np.maximum(1e-12, np.abs(ref))) < 1e-9, col
    print(f"OK: con suavizado = 0 el libro 1 coincide con la referencia sin suavizar (e0 sin suavizar = {calc.tabla(q0).ex[0]:.2f})")
print(f"México 2019, ambos sexos: e0 = {t.ex[0]:.2f}, e65 = {t.ex[65]:.2f}, q0 = {t.qx[0] * 1000:.2f} por mil, l65 = {t.lx[65] / 1e5:.3f}")
