"""Gráficos PNG de la tabla de mortalidad, a partir de los valores calculados del libro de Excel.

Uso:  python3 scripts/03_graficos.py
"""
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.ticker import FuncFormatter
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
LIBRO = RAIZ / "resultados" / "T_Mortalidad_INEGI_Mexico_2019.xlsx"
HOJA = {"H": "Tabla hombres", "M": "Tabla mujeres", "A": "Tabla ambos sexos"}
NOM = {"H": "Hombres", "M": "Mujeres", "A": "Ambos sexos"}
COL = {"H": "#2a78d6", "M": "#eb6834", "A": "#1baf7a"}
INK, INK2, GRID, BG = "#0b0b0b", "#52514e", "#e4e3df", "#fcfcfb"
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "axes.edgecolor": GRID, "axes.labelcolor": INK2,
                     "xtick.color": INK2, "ytick.color": INK2, "text.color": INK, "axes.spines.top": False,
                     "axes.spines.right": False, "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.8,
                     "axes.axisbelow": True, "figure.facecolor": BG, "axes.facecolor": BG, "savefig.facecolor": BG})
wb = load_workbook(LIBRO, data_only=True)
T = {}
for s, h in HOJA.items():
    ws = wb[h]
    T[s] = {k: np.array([ws[f"{c}{5 + i}"].value for i in range(101)], float)
            for c, k in zip("BCDEFGH", ["x", "qx", "lx", "dx", "Lx", "Tx", "ex"])}
x = T["A"]["x"]
SUB = "México, 2019. Defunciones y nacimientos: INEGI (EDR, ENR); población: Censo 2020"


def base(titulo, ylabel, xlim=(0, 100)):
    fig, ax = plt.subplots(figsize=(9, 5))
    ax.set_title(titulo, loc="left", fontsize=13, fontweight="bold", pad=22)
    ax.text(0, 1.02, SUB, transform=ax.transAxes, fontsize=8.5, color=INK2)
    ax.set_xlabel("Edad x (años cumplidos)")
    ax.set_ylabel(ylabel)
    ax.set_xlim(*xlim)
    return fig, ax


def guardar(fig, nombre):
    fig.tight_layout()
    fig.savefig(RAIZ / "resultados" / nombre, dpi=160)
    plt.close(fig)
    print("png", nombre)


def etiquetas_fin(ax, xs, series, fmt, dx=1, dy=None):
    for s, y in series:
        ax.text(xs[-1] + dx, y[-1] if dy is None else dy.get(s, y[-1]), NOM[s], color=INK, fontsize=9, va="center")


# 1. lx
fig, ax = base("Sobrevivientes lx de una generación ficticia de 100,000 nacidos", "lx")
for s in "HMA":
    ax.plot(x, T[s]["lx"], color=COL[s], lw=2.2, label=NOM[s])
ax.axhline(50000, color=INK2, lw=1, ls=(0, (4, 3)))
ax.text(1, 51500, "mitad de la generación", fontsize=8.5, color=INK2)
ax.axvline(65, color=INK2, lw=0.8, ls=(0, (2, 3)))
ax.scatter([65] * 3, [T[s]["lx"][65] for s in "HMA"], color=[COL[s] for s in "HMA"], zorder=3, s=30)
ax.text(20, 22000, "Llegan con vida a los 65 años (de 100,000):\n" + "\n".join(
    f"  {NOM[s]}: {T[s]['lx'][65]:,.0f}" for s in "MAH"), fontsize=9, color=INK, va="center")
ax.set_ylim(0, 102000)
ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{int(v):,}"))
ax.legend(frameon=False, loc="lower left")
guardar(fig, "lx_sobrevivientes.png")

# 2. qx (escala log)
fig, ax = base("Probabilidad de morir qx por edad (escala logarítmica)", "qx (log)", xlim=(0, 99))
for s in "HMA":
    ax.plot(x[:100], T[s]["qx"][:100], color=COL[s], lw=2.2, label=NOM[s])
ax.set_yscale("log")
ax.set_ylim(1e-4, 0.6)
ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:g}"))
ax.grid(True, which="minor", color=GRID, lw=0.4)
ax.legend(frameon=False, loc="upper left")
ax.text(60, 2.2e-4, "A partir de ~90 años qx se aplana: el Censo exagera\nla población de edades avanzadas (ver análisis)", fontsize=8.5, color=INK2)
guardar(fig, "qx_log.png")

# 3. dx
fig, ax = base("Defunciones dx de la tabla, por edad", "dx (de 100,000 nacidos)", xlim=(0, 99))
for s in "HMA":
    ax.plot(x[:100], T[s]["dx"][:100], color=COL[s], lw=2.2, label=NOM[s])
ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{int(v):,}"))
ax.legend(frameon=False, loc="upper left")
ax.annotate("primer año de vida", (0.6, 1000), xytext=(7, 1500), fontsize=9, color=INK,
            arrowprops=dict(arrowstyle="-", color=INK2, lw=0.8))
guardar(fig, "dx_defunciones.png")

# 4. ex
fig, ax = base("Esperanza de vida ex por edad", "ex (años por vivir)")
for s in "MAH":
    ax.plot(x, T[s]["ex"], color=COL[s], lw=2.2, label=f"{NOM[s]} (e0 = {T[s]['ex'][0]:.1f})")
    ax.scatter([0], [T[s]["ex"][0]], color=COL[s], zorder=3, s=36)
ax.set_ylim(0, 85)
ax.legend(frameon=False, loc="upper right")
guardar(fig, "ex_esperanza.png")

# 5. sobremortalidad masculina
r = T["H"]["qx"][1:100] / T["M"]["qx"][1:100]
fig, ax = base("Sobremortalidad masculina: qx de hombres / qx de mujeres", "veces", xlim=(1, 99))
ax.plot(x[1:100], r, color=COL["H"], lw=2.2)
ax.axhline(1, color=INK2, lw=1)
pico = int(np.argmax(r[14:39]) + 15)
ax.annotate(f"{r[pico - 1]:.1f} veces a los {pico} años", (pico, r[pico - 1]), xytext=(pico + 8, r[pico - 1] + 0.15),
            fontsize=9, color=INK, arrowprops=dict(arrowstyle="-", color=INK2, lw=0.8))
ax.set_ylim(0.5, max(r) * 1.1)
guardar(fig, "sobremortalidad_masculina.png")

# 6. e0 vs ONU
wi = wb["Indicadores"]
fila = next(i for i in range(1, 40) if wi.cell(i, 1).value == "e0 con datos de INEGI (esta tabla)")
inegi = [wi.cell(fila, c).value for c in (2, 3, 4)]
onu = [wi.cell(fila + 1, c).value for c in (2, 3, 4)]
fig, ax = base("e0 2019: tabla con datos de INEGI vs. estimación de la ONU", "e0 (años)", xlim=(-0.6, 2.6))
xs = np.arange(3)
ax.bar(xs - 0.2, inegi, 0.38, color="#2a78d6", label="Esta tabla (INEGI)")
ax.bar(xs + 0.2, onu, 0.38, color="#9aa0a6", label="ONU, WPP 2024")
for xi, a, b in zip(xs, inegi, onu):
    ax.text(xi - 0.2, a + 0.4, f"{a:.1f}", ha="center", fontsize=9)
    ax.text(xi + 0.2, b + 0.4, f"{b:.1f}", ha="center", fontsize=9)
    ax.text(xi, max(a, b) + 2.3, f"INEGI − ONU: +{a - b:.1f} años", ha="center", fontsize=9, color=INK2)
ax.set_xticks(xs)
ax.set_xticklabels(["Hombres", "Mujeres", "Ambos sexos"])
ax.set_xlabel("")
ax.set_ylim(60, 85)
ax.legend(frameon=False, loc="upper left")
guardar(fig, "e0_comparacion_onu.png")
