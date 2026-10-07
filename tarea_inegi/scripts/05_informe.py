"""Extrae cifras del libro recalculado y de los datos, y genera el documento Word de análisis (llama a scripts/05_informe.js).

Uso:  python3 scripts/05_informe.py
"""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import calculo_inegi as calc  # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent
wb = load_workbook(RAIZ / "resultados" / "T_Mortalidad_INEGI_Mexico_2019.xlsx", data_only=True)
ws = wb["Tabla de mortalidad"]
T = {k: np.array([ws[f"{c}{5 + i}"].value for i in range(101)], float) for c, k in zip("BCDEFGH", ["x", "qx", "lx", "dx", "Lx", "Tx", "ex"])}
q, l, d, e = T["qx"], T["lx"], T["dx"], T["ex"]
qh, qm = calc.qx("H")[0], calc.qx("M")[0]
razon = qh[1:100] / qm[1:100]                       # edades 1..99
r15_39 = razon[14:39]
ref = pd.read_csv(RAIZ / "datos" / "conapo_e0_2019_referencia.csv").set_index("sexo").e0_2019
d_arch = pd.read_csv(RAIZ / "datos" / "inegi_defunciones_2019_por_archivo.csv")
n_arch = pd.read_csv(RAIZ / "datos" / "inegi_nacimientos_2019_por_archivo.csv")
dfn, nac, pob = calc.cargar()
datos = {
    "e0": e[0], "e1": e[1], "e15": e[15], "e65": e[65], "q0_mil": q[0] * 1000, "5q0_mil": (1 - l[5] / l[0]) * 1000,
    "l65": l[65] / 1e5, "l80": l[80] / 1e5, "l100": l[100] / 1e5,
    "mediana": int(T["x"][np.where(l >= 5e4)[0][-1]]), "moda": int(15 + np.argmax(d[15:100])),
    "edad_supera_q0": int(1 + np.argmax(q[1:100] > q[0])), "q_min": float(q[1:30].min()), "edad_q_min": int(1 + np.argmin(q[1:30])), "q50": q[50], "q80": q[80], "q99": q[99],
    "def_inf": float(ws["L7"].value), "nacimientos": float(ws["L6"].value),
    "def_total": int(d_arch.defunciones_ocurridas_2019.sum()), "nac_arch": [int(x) for x in (n_arch.hombres + n_arch.mujeres + n_arch.sexo_no_especificado)],
    "def_arch": [int(x) for x in d_arch.defunciones_ocurridas_2019], "pob_total": int(pob.values.sum()),
    "e0_h": calc.e0_continua(qh), "e0_m": calc.e0_continua(qm), "e0_a": calc.e0_continua(calc.qx("A")[0]),
    "ref_h": float(ref["Hombres"]), "ref_m": float(ref["Mujeres"]), "ref_a": float(ref["Ambos sexos"]),
    "razon_media": float(r15_39.mean()), "razon_max": float(r15_39.max()), "razon_edad": int(15 + np.argmax(r15_39)),
    "extracto": [[int(x), T["qx"][x], T["lx"][x], T["dx"][x], T["Lx"][x], T["Tx"][x], T["ex"][x]] for x in (0, 1, 5, 10, 20, 30, 40, 50, 60, 65, 70, 80, 90, 99, 100)],
}
salida = RAIZ / "resultados" / "Analisis_Tabla_Mortalidad_INEGI_2019.docx"
with tempfile.TemporaryDirectory() as tmp:
    ruta = Path(tmp) / "datos.json"
    ruta.write_text(json.dumps(datos, ensure_ascii=False))
    subprocess.run(["node", str(RAIZ / "scripts" / "05_informe.js"), str(ruta), str(RAIZ / "resultados"), str(salida)], check=True)
print("informe", salida)
