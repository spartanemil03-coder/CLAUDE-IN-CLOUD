"""Descarga datos abiertos de INEGI y deja en datos/ solo los conteos necesarios para la tabla de mortalidad.

Fuentes (todas INEGI):
  - Estadísticas de Defunciones Registradas (EDR), conjunto de datos abiertos 2019, 2020 y 2021 (CSV):
    https://www.inegi.org.mx/programas/edr/#datos_abiertos
  - Estadística de Nacimientos Registrados (ENR / natalidad), conjunto de datos abiertos 2019, 2020 y 2021 (CSV):
    https://www.inegi.org.mx/programas/natalidad/#datos_abiertos
  - Censo de Población y Vivienda 2020, tabulados del cuestionario básico, "Población 3" (edad desplegada y sexo):
    https://www.inegi.org.mx/programas/ccpv/2020/#tabulados

Año de referencia: 2019 (último año completo antes de la pandemia). Defunciones y nacimientos se cuentan por AÑO DE
OCURRENCIA = 2019, tomando los registrados en 2019 y en los dos años siguientes (registro tardío).

Uso:  python3 scripts/01_descargar_inegi.py [carpeta_datos] [carpeta_temporal]
"""
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
DATOS = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / "datos"
TMP = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(tempfile.mkdtemp(prefix="inegi_"))
ANIO = 2019
BASE = "https://www.inegi.org.mx/contenidos/programas"
URL_EDR = BASE + "/edr/datosabiertos/defunciones/{y}/conjunto_de_datos_defunciones_registradas_{y}_csv.zip"
URL_ENR = BASE + "/natalidad/datosabiertos/{y}/conjunto_de_datos_natalidad_{y}_csv.zip"
URL_CENSO = BASE + "/ccpv/2020/tabulados/cpv2020_b_eum_01_poblacion.xlsx"


def bajar(url, destino):
    destino.parent.mkdir(parents=True, exist_ok=True)
    if not destino.exists():
        subprocess.run(["curl", "-sS", "-f", "-L", "-m", "900", "-o", str(destino), url], check=True)
    return destino


def csv_de_zip(zpath):
    """Extrae el CSV de la carpeta conjunto_de_datos del zip (ignora Nota.txt y bitácoras)."""
    out = zpath.with_suffix("")
    with zipfile.ZipFile(zpath) as z:
        nombres = [n for n in z.namelist() if "conjunto_de_datos/" in n and n.lower().endswith(".csv")
                   and "bitacora" not in n.lower()]
        assert len(nombres) == 1, nombres
        z.extract(nombres[0], out)
    return out / nombres[0]


def leer(csv, cols):
    d = pd.read_csv(csv, encoding="latin1", usecols=lambda c: c.lower() in cols, low_memory=False)
    d.columns = d.columns.str.lower()
    return d


def defunciones():
    partes = []
    for y in (ANIO, ANIO + 1, ANIO + 2):
        z = bajar(URL_EDR.format(y=y), TMP / f"edr{y}.zip")
        d = leer(csv_de_zip(z), {"sexo", "edad", "anio_ocur"})
        partes.append(d[d.anio_ocur == ANIO])
        print("EDR", y, "defunciones ocurridas en", ANIO, ":", len(partes[-1]))
    d = pd.concat(partes, ignore_index=True)
    # EDAD: 1001-3999 = menores de un año (horas, días, meses); 4001-4120 = años cumplidos; 4998 = no especificada
    edad = np.where(d.edad < 4000, 0, np.where(d.edad == 4998, -1, d.edad - 4000))
    d["edad"] = np.where(edad < 0, -1, np.minimum(edad, 100))
    tab = pd.crosstab(d.edad, d.sexo).reindex(columns=[1, 2, 9], fill_value=0)
    tab.columns = ["hombres", "mujeres", "sexo_no_especificado"]
    tab = tab.reindex(list(range(101)) + [-1], fill_value=0)
    tab.index = [str(i) if i >= 0 else "NE" for i in tab.index]
    tab.index.name = "edad"
    return tab


def nacimientos():
    partes = []
    for y in (ANIO, ANIO + 1, ANIO + 2):
        z = bajar(URL_ENR.format(y=y), TMP / f"enr{y}.zip")
        d = leer(csv_de_zip(z), {"sexo", "ano_nac"})
        partes.append(d[d.ano_nac == ANIO])
        print("ENR", y, "nacimientos ocurridos en", ANIO, ":", len(partes[-1]))
    d = pd.concat(partes, ignore_index=True)
    c = d.sexo.value_counts().reindex([1, 2, 9], fill_value=0)
    return pd.DataFrame({"hombres": [c[1]], "mujeres": [c[2]], "sexo_no_especificado": [c[9]]})


def poblacion_censo():
    x = bajar(URL_CENSO, TMP / "cpv2020_b_eum_01_poblacion.xlsx")
    ws = load_workbook(x, read_only=True, data_only=True)["03"]
    filas = []
    for r in ws.iter_rows(min_row=1, values_only=True):
        if r[0] == "Estados Unidos Mexicanos" and r[1] not in ("Total",):
            filas.append(r[1:5])
    edades = [str(i) for i in range(101)] + ["NE"]
    assert len(filas) == 102, len(filas)           # 0..99, 100 y más, No especificado
    df = pd.DataFrame({"edad": edades, "hombres": [f[2] for f in filas], "mujeres": [f[3] for f in filas]})
    assert filas[0][0].startswith("00") and filas[100][0].startswith("100") and filas[101][0].startswith("No esp")
    return df.set_index("edad")


def main():
    DATOS.mkdir(parents=True, exist_ok=True)
    d = defunciones()
    d.to_csv(DATOS / f"inegi_defunciones_{ANIO}_por_edad_sexo.csv")
    n = nacimientos()
    n.to_csv(DATOS / f"inegi_nacimientos_{ANIO}_por_sexo.csv", index=False)
    p = poblacion_censo()
    p.to_csv(DATOS / "inegi_censo2020_poblacion_por_edad_sexo.csv")
    print("defunciones", int(d.values.sum()), "| nacimientos", int(n.values.sum()), "| población", int(p.values.sum()))


if __name__ == "__main__":
    main()
