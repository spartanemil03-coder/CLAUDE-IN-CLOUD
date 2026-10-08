"""Reto Actinver 2026: cribado de las 145 acciones y comparación de portafolios candidatos.

Lee actinver/datos/Reto_Actinver_145_acciones.xlsx (hojas Acciones y Covarianza) y escribe en actinver/resultados/.
Todo es descriptivo del pasado: no predice rendimientos. Las probabilidades suponen rendimiento esperado 0.
"""
import warnings
from pathlib import Path
from statistics import NormalDist

import numpy as np
import openpyxl
import pandas as pd

warnings.filterwarnings("ignore")
RAIZ = Path(__file__).resolve().parent.parent
XLSX = RAIZ / "datos" / "Reto_Actinver_145_acciones.xlsx"
SALIDA = RAIZ / "resultados"

CAPITAL = 1_000_000
COMISION = 0.001 * 1.16       # 0.10% + IVA, por operación
SESIONES = 25                 # sesiones que faltan (misma cifra que la hoja Mi portafolio)

COLS = ["emisora", "tipo", "ultimo", "var_dia", "p_compra", "p_venta", "spread", "vol_compra", "vol_venta",
        "operable", "r1m", "r3m", "r6m", "r12m", "vol_anual", "mov_1s", "peor_caida", "beta", "mejor_vent",
        "peor_vent", "perfil"]

PORTAFOLIOS = {
    "A Prudente": {"GFNORTE O": .20, "GMEXICO B": .20, "AC *": .20, "FUNO 11": .20, "MSFT *": .20},
    "B Tu idea + diversificadores": {"NVDA *": .30, "GFNORTE O": .20, "GMEXICO B": .20, "MSFT *": .15, "GOOGL *": .15},
    "C Ranking (alto riesgo)": {"NVDA *": .40, "TSLA *": .20, "AVGO *": .15, "GMEXICO B": .15, "GFNORTE O": .10},
}


def cargar():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["Acciones"]
    df = pd.DataFrame([[c.value for c in ws[r]] for r in range(5, 150)], columns=COLS)
    cs = wb["Covarianza"]
    nombres = [c.value for c in cs[4][1:145]]
    cov = np.array([[cs.cell(6 + i, 2 + j).value or 0 for j in range(144)] for i in range(144)], dtype=float)
    return df, pd.DataFrame(cov, index=nombres, columns=nombres)


def cribado(df):
    d = df.copy()
    d["costo_ida_vuelta"] = d.spread + 2 * COMISION
    d["liq_mejor_nivel_mxn"] = np.minimum(d.vol_compra, d.vol_venta) * d.ultimo
    d["ejecucion"] = np.where(d.operable != "Sí", "sin contraparte",
                      np.where(d.spread <= 0.006, "barata (<=0.6%)",
                      np.where(d.spread <= 0.02, "media (0.6-2%)", "cara (>2%)")))
    cols = ["emisora", "tipo", "ejecucion", "spread", "costo_ida_vuelta", "liq_mejor_nivel_mxn", "vol_anual",
            "mov_1s", "beta", "r1m", "r3m", "r6m", "r12m", "peor_caida", "mejor_vent", "peor_vent"]
    return d[cols].sort_values(["ejecucion", "vol_anual"])


def evalua(nombre, pesos, df, cov):
    assert len(pesos) >= 5 and max(pesos.values()) <= 0.5 + 1e-9, nombre
    assert abs(sum(pesos.values()) - 1) < 1e-9, nombre
    d = df.set_index("emisora")
    w = pd.Series(pesos)
    sig_anual = float(np.sqrt(w @ cov.loc[w.index, w.index] @ w))
    sig = sig_anual * np.sqrt(SESIONES / 252)
    costo = float(sum(CAPITAL * p * (d.loc[e, "spread"] + 2 * COMISION) for e, p in pesos.items()))
    beta = float(sum(p * d.loc[e, "beta"] for e, p in pesos.items()))
    # Ganancia neta ~ Normal(-costo, (sig*CAPITAL)^2); media de rendimiento 0.
    prob = lambda x: 1 - NormalDist(-costo, sig * CAPITAL).cdf(x)
    return {"portafolio": nombre, "emisoras": ", ".join(f"{e} {p:.0%}" for e, p in pesos.items()),
            "vol_anual": sig_anual, "mov_1sigma_pct": sig, "mov_1sigma_mxn": sig * CAPITAL,
            "costo_estimado_mxn": costo, "costo_pct": costo / CAPITAL, "beta_vs_IPC": beta,
            "P_ganar_algo": prob(0), "P_ganar_100k": prob(100_000), "P_ganar_200k": prob(200_000),
            "P_perder_100k": 1 - prob(-100_000)}


def ordenes(nombre, pesos, df):
    """Acciones enteras a comprar al precio de venta, dejando ~0.3% de efectivo para la comisión."""
    d = df.set_index("emisora")
    filas = []
    for e, p in pesos.items():
        titulos = int(p * CAPITAL * 0.997 // d.loc[e, "p_venta"])
        filas.append({"portafolio": nombre, "emisora": e, "peso": p, "precio_venta": d.loc[e, "p_venta"],
                      "titulos": titulos, "importe": round(titulos * d.loc[e, "p_venta"], 2)})
    return filas


def main():
    df, cov = cargar()
    SALIDA.mkdir(exist_ok=True)
    c = cribado(df)
    c.round(5).to_csv(SALIDA / "cribado_145_acciones.csv", index=False)
    p = pd.DataFrame([evalua(n, w, df, cov) for n, w in PORTAFOLIOS.items()])
    p.round(4).to_csv(SALIDA / "portafolios_candidatos.csv", index=False)
    o = pd.DataFrame([f for n, w in PORTAFOLIOS.items() for f in ordenes(n, w, df)])
    o.to_csv(SALIDA / "ordenes_sugeridas.csv", index=False)
    pd.set_option("display.width", 250, "display.max_columns", 30)
    print(c.ejecucion.value_counts().to_string())
    print(p.drop(columns="emisoras").round(3).to_string(index=False))
    sel = ["NVDA *", "GFNORTE O"]
    print(o.to_string(index=False))
    print("\nCorrelación NVDA vs GFNORTE:",
          round(cov.loc[sel[0], sel[1]] / np.sqrt(cov.loc[sel[0], sel[0]] * cov.loc[sel[1], sel[1]]), 2))


if __name__ == "__main__":
    main()
