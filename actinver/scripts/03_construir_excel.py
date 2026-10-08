"""Reto Actinver 2026: agrega al Excel original las hojas Pronóstico, Portafolios y Beta, y llena Mi portafolio con el ganador.

Entrada: datos/Reto_Actinver_145_acciones.xlsx (no se modifica) y resultados/*.csv de 02_escenarios.py.
Salida:  resultados/Reto_Actinver_mejorado.xlsx (después hay que recalcular con recalc.py de la skill xlsx).
"""
import warnings
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

warnings.filterwarnings("ignore")
RAIZ = Path(__file__).resolve().parent.parent
ORIGINAL = RAIZ / "datos" / "Reto_Actinver_145_acciones.xlsx"
SALIDA = RAIZ / "resultados" / "Reto_Actinver_mejorado.xlsx"

AZUL = "1F4E78"
F_TIT = Font(name="Calibri", size=14, bold=True, color=AZUL)
F_SEC = Font(name="Calibri", size=12, bold=True, color=AZUL)
F_ENC = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
F_TXT = Font(name="Calibri", size=10)
F_NEG = Font(name="Calibri", size=10, bold=True)
F_NOTA = Font(name="Calibri", size=9)
F_FIJO = Font(name="Calibri", size=10, italic=True, color="595959")
REL_ENC = PatternFill("solid", fgColor=AZUL)
REL_IN = PatternFill("solid", fgColor="FFF2CC")
REL_GAN = PatternFill("solid", fgColor="E2EFDA")
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
FIN = Side(style="thin", color="BFBFBF")
BORDE = Border(bottom=FIN)
PESOS = '#,##0;[Red]-#,##0;"-"'
PCT = '0.0%;[Red]-0.0%;"-"'


def enc(ws, fila, textos, col0=1):
    for j, t in enumerate(textos):
        c = ws.cell(fila, col0 + j, t)
        c.font, c.fill, c.alignment = F_ENC, REL_ENC, CENTRO


def celda(ws, ref, valor, fmt=None, font=F_TXT, fill=None, centro=True):
    c = ws[ref]
    c.value, c.font = valor, font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if centro:
        c.alignment = Alignment(horizontal="center")
    return c


def main():
    res = pd.read_csv(RAIZ / "resultados" / "pronostico_escenarios.csv")
    pesos = pd.read_csv(RAIZ / "resultados" / "pesos_portafolios.csv")
    carteras = list(res.portafolio)
    ganador = res.loc[res.ganancia_por_riesgo.idxmax(), "portafolio"]

    wb = load_workbook(ORIGINAL)
    cov = wb["Covarianza"]
    nombres = [c.value for c in cov[4][1:145]]
    n = len(nombres)
    CAP = "'Mi portafolio'!$C$4"
    COM = "'Mi portafolio'!$C$5"
    SES = "'Mi portafolio'!$H$4"
    rango_cov = f"Covarianza!$B$6:${L(1 + n)}${5 + n}"

    # ---------------- Beta (soporte) ----------------
    be = wb.create_sheet("Beta")
    be.sheet_view.showGridLines = False
    be["A1"], be["A1"].font = "Beta de cada emisora contra la canasta de las 144 acciones (mismo peso)", F_TIT
    be["A2"], be["A2"].font = ("Sale de la hoja Covarianza: beta = cov(emisora, canasta) / var(canasta). Costo ida y vuelta = "
                                "diferencia compra-venta + 2 comisiones. Hoja de apoyo.", F_NOTA)
    enc(be, 4, ["Emisora", "Beta vs canasta", "Costo ida y vuelta", "¿Operable hoy?"])
    for i, e in enumerate(nombres):
        r, cr = 5 + i, 6 + i
        be.cell(r, 1, e).font = F_TXT
        c = be.cell(r, 2, f"=SUM(Covarianza!$B{cr}:${L(1 + n)}{cr})/{n}/(SUM({rango_cov})/{n}^2)")
        c.number_format, c.font, c.alignment = "0.00", F_TXT, Alignment(horizontal="center")
        m = f"MATCH($A{r},Acciones!$A$5:$A$148,0)"
        c = be.cell(r, 3, f"=IF(ISNUMBER(INDEX(Acciones!$G$5:$G$148,{m})),INDEX(Acciones!$G$5:$G$148,{m}),1)+2*{COM}")
        c.number_format, c.font, c.alignment = "0.00%", F_TXT, Alignment(horizontal="center")
        c = be.cell(r, 4, f"=INDEX(Acciones!$J$5:$J$148,{m})")
        c.font, c.alignment = F_TXT, Alignment(horizontal="center")
    for col, w in zip("ABCD", (14, 14, 16, 16)):
        be.column_dimensions[col].width = w
    be.freeze_panes = "B5"
    rbe = f"Beta!$A$5:$A${4 + n}"
    bbe, cbe = f"Beta!$B$5:$B${4 + n}", f"Beta!$C$5:$C${4 + n}"

    # ---------------- Portafolios ----------------
    po = wb.create_sheet("Portafolios")
    po.sheet_view.showGridLines = False
    po["A1"], po["A1"].font = "Portafolios candidatos: pesos y ganancia por escenario", F_TIT
    po["A2"], po["A2"].font = ("Los pesos (celdas amarillas) se pueden cambiar; las métricas se recalculan. "
                                "A, B y C son los de la propuesta anterior; D y E salen del optimizador.", F_NOTA)
    enc(po, 4, ["Emisora", "Beta vs canasta", "Costo ida y vuelta"] + carteras)
    po.row_dimensions[4].height = 42
    union = list(dict.fromkeys(pesos.emisora))
    p0, p1 = 5, 4 + len(union)
    mapa = {(r.portafolio, r.emisora): r.peso for r in pesos.itertuples()}
    for i, e in enumerate(union):
        r = p0 + i
        po.cell(r, 1, e).font = F_TXT
        c = po.cell(r, 2, f"=INDEX({bbe},MATCH($A{r},{rbe},0))")
        c.number_format, c.font, c.alignment = "0.00", F_TXT, Alignment(horizontal="center")
        c = po.cell(r, 3, f"=INDEX({cbe},MATCH($A{r},{rbe},0))")
        c.number_format, c.font, c.alignment = "0.00%", F_TXT, Alignment(horizontal="center")
        for j, k in enumerate(carteras):
            c = po.cell(r, 4 + j, mapa.get((k, e), 0))
            c.number_format, c.font, c.fill, c.alignment = PCT, F_TXT, REL_IN, Alignment(horizontal="center")
    m0 = p1 + 2
    etiquetas = ["Suma de pesos", "Emisoras distintas (mín. 5)", "Peso máximo (máx. 50%)", "Revisión de reglas",
                 "Beta vs canasta", "Costo estimado ida y vuelta ($)", "Ganancia neta, escenario bajista ($)",
                 "Ganancia neta, escenario base ($)", "Ganancia neta, escenario alcista ($)",
                 "Ganancia esperada (ponderada por probabilidad, $)"]
    for k, t in enumerate(etiquetas):
        po.cell(m0 + k, 1, t).font = F_NEG
    for j in range(len(carteras)):
        col = L(4 + j)
        w = f"{col}${p0}:{col}${p1}"
        fm = {
            0: (f"=SUM({w})", PCT),
            1: (f'=COUNTIF({w},">0")', "0"),
            2: (f"=MAX({w})", PCT),
            3: (f'=IF(AND(ABS({col}{m0}-1)<0.0005,{col}{m0 + 1}>=5,{col}{m0 + 2}<=0.5),"OK","REVISAR")', "General"),
            4: (f"=SUMPRODUCT({w},$B${p0}:$B${p1})", "0.00"),
            5: (f"=SUMPRODUCT({w},$C${p0}:$C${p1})*{CAP}", PESOS),
            6: (f"={CAP}*{col}{m0 + 4}*'Pronóstico'!$C$8-{col}{m0 + 5}", PESOS),
            7: (f"={CAP}*{col}{m0 + 4}*'Pronóstico'!$C$9-{col}{m0 + 5}", PESOS),
            8: (f"={CAP}*{col}{m0 + 4}*'Pronóstico'!$C$10-{col}{m0 + 5}", PESOS),
            9: (f"='Pronóstico'!$B$8*{col}{m0 + 6}+'Pronóstico'!$B$9*{col}{m0 + 7}+'Pronóstico'!$B$10*{col}{m0 + 8}", PESOS),
        }
        for k, (f, fmt) in fm.items():
            celda(po, f"{col}{m0 + k}", f, fmt, F_NEG)
    po.column_dimensions["A"].width = 46
    for col in "BC":
        po.column_dimensions[col].width = 14
    for j in range(len(carteras)):
        po.column_dimensions[L(4 + j)].width = 18
    po.freeze_panes = "D5"

    # ---------------- Pronóstico ----------------
    pr = wb.create_sheet("Pronóstico")
    pr.sheet_view.showGridLines = False
    pr["A1"], pr["A1"].font = "Pronóstico por escenarios y ganador", F_TIT
    pr["A2"], pr["A2"].font = ("Hipotético y descriptivo del pasado: no predice qué acción sube más. Cada escenario supone cómo se mueve "
                                "el mercado (canasta de las 144 acciones) y cada acción lo sigue según su beta.", F_NOTA)
    pr["A4"], pr["A4"].font = "Supuestos (celdas amarillas editables)", F_SEC
    celda(pr, "A5", "Mov. típico (±1σ) de la canasta en las sesiones que faltan", None, F_NEG, centro=False)
    celda(pr, "B5", f"=SQRT(SUM({rango_cov})/{n}^2*{SES}/252)", "0.0%", F_NEG)
    enc(pr, 7, ["Escenario", "Probabilidad", "Rend. de la canasta", "Cómo se lee"])
    filas = [("Bajista", 0.25, "=-$B$5", "El mercado baja 1σ"),
             ("Base", 0.50, 0.01, "Sube ~1% en el lapso (≈10% anual)"),
             ("Alcista", 0.25, "=$B$5", "El mercado sube 1σ")]
    for k, (nom, p, rr, txt) in enumerate(filas):
        r = 8 + k
        celda(pr, f"A{r}", nom, None, F_NEG, centro=False)
        celda(pr, f"B{r}", p, "0%", F_TXT, REL_IN)
        celda(pr, f"C{r}", rr, PCT, F_TXT, REL_IN)
        celda(pr, f"D{r}", txt, None, F_NOTA, centro=False)
    celda(pr, "A11", "Suma de probabilidades (debe ser 100%)", None, F_NOTA, centro=False)
    celda(pr, "B11", "=SUM(B8:B10)", "0%", F_NOTA)
    celda(pr, "A12", "Rend. medio esperado de la canasta", None, F_NOTA, centro=False)
    celda(pr, "B12", "=SUMPRODUCT(B8:B10,C8:C10)", "0.0%", F_NOTA)

    pr["A14"], pr["A14"].font = "Resultado por portafolio (pesos sobre el capital, ya con costos de operar)", F_SEC
    cab = ["Portafolio", "Emisoras", "Beta", "Costo estimado ($)", "Bajista ($)", "Base ($)", "Alcista ($)",
           "Ganancia esperada ($)", "Vol. anual*", "Desv. de la ganancia ($)*", "Ganancia por riesgo",
           "Peor 5% ($)*", "Mejor 5% ($)*", "Prob. de ganar*"]
    enc(pr, 15, cab)
    pr.row_dimensions[15].height = 42
    for j, k in enumerate(carteras):
        r, col = 16 + j, L(4 + j)
        x = res.iloc[j]
        celda(pr, f"A{r}", k, None, F_NEG, REL_GAN if k == ganador else None, centro=False)
        celda(pr, f"B{r}", f"=Portafolios!{col}{m0 + 1}", "0")
        celda(pr, f"C{r}", f"=Portafolios!{col}{m0 + 4}", "0.00")
        celda(pr, f"D{r}", f"=Portafolios!{col}{m0 + 5}", PESOS)
        for off, cc in zip((6, 7, 8, 9), "EFGH"):
            celda(pr, f"{cc}{r}", f"=Portafolios!{col}{m0 + off}", PESOS, F_NEG if cc == "H" else F_TXT)
        celda(pr, f"I{r}", round(float(x.vol_anual), 4), "0.0%", F_FIJO)
        celda(pr, f"J{r}", round(float(x.desv_ganancia), 0), PESOS, F_FIJO)
        celda(pr, f"K{r}", f"=IF(J{r}>0,H{r}/J{r},0)", "0.000", F_NEG)
        celda(pr, f"L{r}", round(float(x.P5), 0), PESOS, F_FIJO)
        celda(pr, f"M{r}", round(float(x.P95), 0), PESOS, F_FIJO)
        celda(pr, f"N{r}", round(float(x.P_ganar), 3), "0.0%", F_FIJO)
    ult = 15 + len(carteras)
    pr["A22"], pr["A22"].font = "Ganadores", F_SEC
    celda(pr, "A23", "Ganador: mayor ganancia esperada por riesgo", None, F_NEG, centro=False)
    celda(pr, "B23", f"=INDEX($A$16:$A${ult},MATCH(MAX($K$16:$K${ult}),$K$16:$K${ult},0))", None, F_NEG, REL_GAN, centro=False)
    for k, (t, cc) in enumerate([("Gana si el mercado baja", "E"), ("Gana en el escenario base", "F"),
                                 ("Gana si el mercado sube", "G"), ("Mayor ganancia esperada", "H")]):
        celda(pr, f"A{24 + k}", t, None, F_TXT, centro=False)
        celda(pr, f"B{24 + k}", f"=INDEX($A$16:$A${ult},MATCH(MAX({cc}$16:{cc}${ult}),{cc}$16:{cc}${ult},0))", None, F_TXT, centro=False)

    notas = [
        "Cómo se calcula: ganancia = capital × (beta del portafolio × rendimiento de la canasta en el escenario) − costo de operar. El costo supone ida y vuelta completa (diferencia compra-venta + 2 comisiones); si se valúa a precio de mercado sin vender, es menor.",
        "No hay pronóstico por emisora: ninguna acción tiene ventaja propia en el modelo. Lo que cambia entre portafolios es cuánto mercado toman (beta), cuánto riesgo propio cargan y cuánto cuesta operarlos.",
        "* Columnas grises (cursiva): simulación de 40,000 trayectorias en Python con la covarianza completa del último año y media = beta × rendimiento medio de la canasta (semilla 2026). El riesgo coincide con el de la hoja Mi portafolio. No se recalculan si cambias los supuestos de arriba; las demás sí.",
        "'Ganancia por riesgo' = ganancia esperada ÷ desviación de la ganancia simulada. Los portafolios D y E se optimizaron con los mismos datos del último año, así que su ventaja sobre A, B y C es optimista.",
        "Las betas, volatilidades y spreads son de una foto del 8 de octubre. Verifica precios, volumen y fechas de resultados trimestrales antes de operar. No es una recomendación de inversión.",
    ]
    for k, t in enumerate(notas):
        c = pr.cell(29 + k, 1, t)
        c.font, c.alignment = F_NOTA, Alignment(wrap_text=True, vertical="top")
        pr.merge_cells(start_row=29 + k, start_column=1, end_row=29 + k, end_column=14)
        pr.row_dimensions[29 + k].height = 26
    pr.column_dimensions["A"].width = 48
    for j in range(2, 15):
        pr.column_dimensions[L(j)].width = 13

    ch = BarChart()
    ch.type, ch.grouping = "col", "clustered"
    ch.title = "Ganancia neta por escenario ($)"
    ch.add_data(Reference(pr, min_col=5, max_col=7, min_row=15, max_row=ult), titles_from_data=True)
    ch.set_categories(Reference(pr, min_col=1, min_row=16, max_row=ult))
    ch.height, ch.width = 9, 24
    ch.y_axis.number_format = "#,##0"
    ch.y_axis.delete = False
    ch.x_axis.delete = False
    ch.x_axis.tickLblPos = "low"
    pr.add_chart(ch, "A35")

    # ---------------- Mi portafolio con el ganador ----------------
    mp = wb["Mi portafolio"]
    for i, (e, w) in enumerate(pesos[pesos.portafolio == ganador][["emisora", "peso"]].itertuples(index=False)):
        mp.cell(8 + i, 1, e)
        mp.cell(8 + i, 2, w)
    # openpyxl descarta la validación de lista de la extensión x14; se vuelve a crear como validación estándar.
    dv = DataValidation(type="list", formula1="Acciones!$A$5:$A$148", allow_blank=True, showErrorMessage=True,
                        errorTitle="Emisora", error="Elige una emisora de la lista (hoja Acciones).")
    dv.add("A8:A19")
    mp.add_data_validation(dv)

    # ---------------- Leeme ----------------
    lm = wb["Leeme"]
    extra = [("Hojas agregadas (versión mejorada)", F_SEC),
             ("• Pronóstico: tres escenarios de mercado (bajista, base, alcista), ganancia en pesos de cada portafolio candidato y ganador.", F_TXT),
             ("• Portafolios: pesos de los 5 candidatos (A a E) y su ganancia por escenario; los pesos se pueden cambiar.", F_TXT),
             ("• Beta: sensibilidad de cada emisora a la canasta de las 144 acciones y su costo de operar (hoja de apoyo).", F_TXT),
             (f"• Mi portafolio viene llenado con el ganador del análisis ({ganador}). Cámbialo si quieres probar otra combinación. No es una recomendación de inversión.", F_TXT)]
    for k, (t, f) in enumerate(extra):
        c = lm.cell(30 + k, 2, t)
        c.font, c.alignment = f, Alignment(wrap_text=True, vertical="top")

    orden = ["Leeme", "Mi portafolio", "Pronóstico", "Portafolios", "Acciones", "Correlaciones", "Covarianza", "Beta"]
    wb._sheets = [wb[s] for s in orden]
    wb.active = 2
    for s in wb:
        s.sheet_view.tabSelected = s.title == "Pronóstico"
    wb.save(SALIDA)
    print("Guardado", SALIDA, "| ganador:", ganador)


if __name__ == "__main__":
    main()
