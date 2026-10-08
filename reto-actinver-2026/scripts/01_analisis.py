"""Reto Actinver 2026: lectura del Excel y cribado de las 145 acciones por costo de operar.

Lee datos/Reto_Actinver_145_acciones.xlsx (hojas Acciones y Covarianza) y escribe en resultados/.
Todo es descriptivo del pasado: no predice rendimientos. Los escenarios y portafolios están en 02_escenarios.py.
"""
import warnings
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

warnings.filterwarnings("ignore")
RAIZ = Path(__file__).resolve().parent.parent
XLSX = RAIZ / "datos" / "Reto_Actinver_145_acciones.xlsx"
SALIDA = RAIZ / "resultados"

CAPITAL = 1_000_000
COMISION = 0.001 * 1.16       # 0.10% + IVA, por operación
SESIONES = 25                 # sesiones que faltan (misma cifra que la hoja Mi portafolio)

COLS = ["emisora", "tipo", "ultimo", "var_dia", "p_compra", "p_venta", "spread", "vol_compra", "vol_venta",
        "operable", "r1m", "r3m", "r6m", "r12m", "vol_anual", "mov_1s", "peor_caida", "beta", "mejor_vent",
        "peor_vent", "perfil"]

PORTAFOLIOS = {
    "A Prudente": {"GFNORTE O": .20, "GMEXICO B": .20, "AC *": .20, "FUNO 11": .20, "MSFT *": .20},
    "B Tu idea + diversificadores": {"NVDA *": .30, "GFNORTE O": .20, "GMEXICO B": .20, "MSFT *": .15, "GOOGL *": .15},
    "C Ranking (alto riesgo)": {"NVDA *": .40, "TSLA *": .20, "AVGO *": .15, "GMEXICO B": .15, "GFNORTE O": .10},
}


def cargar():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["Acciones"]
    df = pd.DataFrame([[c.value for c in ws[r]] for r in range(5, 150)], columns=COLS)
    cs = wb["Covarianza"]
    nombres = [c.value for c in cs[4][1:145]]
    cov = np.array([[cs.cell(6 + i, 2 + j).value or 0 for j in range(144)] for i in range(144)], dtype=float)
    return df, pd.DataFrame(cov, index=nombres, columns=nombres)


def cribado(df):
    d = df.copy()
    d["costo_ida_vuelta"] = d.spread + 2 * COMISION
    d["liq_mejor_nivel_mxn"] = np.minimum(d.vol_compra, d.vol_venta) * d.ultimo
    d["ejecucion"] = np.where(d.operable != "Sí", "sin contraparte",
                      np.where(d.spread <= 0.006, "barata (<=0.6%)",
                      np.where(d.spread <= 0.02, "media (0.6-2%)", "cara (>2%)")))
    cols = ["emisora", "tipo", "ejecucion", "spread", "costo_ida_vuelta", "liq_mejor_nivel_mxn", "vol_anual",
            "mov_1s", "beta", "r1m", "r3m", "r6m", "r12m", "peor_caida", "mejor_vent", "peor_vent"]
    return d[cols].sort_values(["ejecucion", "vol_anual"])


def main():
    df, _ = cargar()
    SALIDA.mkdir(exist_ok=True)
    c = cribado(df)
    c.round(5).to_csv(SALIDA / "cribado_145_acciones.csv", index=False)
    print(c.ejecucion.value_counts().to_string())


if __name__ == "__main__":
    main()
