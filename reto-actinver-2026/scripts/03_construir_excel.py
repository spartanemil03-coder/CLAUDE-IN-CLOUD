"""Reto Actinver 2026: actualiza el Excel original con precios de Yahoo, pronóstico, portafolios de 5 a 12 emisoras y órdenes.

Entrada: datos/Reto_Actinver_145_acciones.xlsx (no se modifica), datos/yahoo/*, resultados/*.csv de 02_pronostico.py.
Salida:  resultados/Reto_Actinver_mejorado.xlsx (después hay que recalcular con recalc.py de la skill xlsx).
"""
import warnings
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

warnings.filterwarnings("ignore")
RAIZ = Path(__file__).resolve().parent.parent
ORIGINAL = RAIZ / "datos" / "Reto_Actinver_145_acciones.xlsx"
RES = RAIZ / "resultados"
SALIDA = RES / "Reto_Actinver_mejorado.xlsx"
HORA = "8 de octubre de 2026, ~11:10 h CDMX"

AZUL = "1F4E78"
F_TIT = Font(name="Calibri", size=14, bold=True, color=AZUL)
F_SEC = Font(name="Calibri", size=12, bold=True, color=AZUL)
F_ENC = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
F_TXT = Font(name="Calibri", size=10)
F_NEG = Font(name="Calibri", size=10, bold=True)
F_NOTA = Font(name="Calibri", size=9)
F_FIJO = Font(name="Calibri", size=10, italic=True, color="595959")
REL_ENC = PatternFill("solid", fgColor=AZUL)
REL_IN = PatternFill("solid", fgColor="FFF2CC")
REL_GAN = PatternFill("solid", fgColor="E2EFDA")
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
PESOS = '#,##0;[Red]-#,##0;"-"'
PCT = '0.0%;[Red]-0.0%;"-"'


def enc(ws, fila, textos, col0=1):
    for j, t in enumerate(textos):
        c = ws.cell(fila, col0 + j, t)
        c.font, c.fill, c.alignment = F_ENC, REL_ENC, CENTRO


def put(ws, ref, valor, fmt=None, font=F_TXT, fill=None, centro=True):
    c = ws[ref]
    c.value, c.font = valor, font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    c.alignment = Alignment(horizontal="center") if centro else Alignment(horizontal="left")
    return c


def titulo(ws, t, sub):
    ws.sheet_view.showGridLines = False
    ws["A1"].value, ws["A1"].font = t, F_TIT
    ws["A2"].value, ws["A2"].font = sub, F_NOTA


def valor(v):
    if isinstance(v, (bool, np.bool_)):
        return "Sí" if v else "No"
    if isinstance(v, str):
        return v
    return None if pd.isna(v) else float(v)


def main():
    e = pd.read_csv(RES / "pronostico_emisoras.csv")
    res = pd.read_csv(RES / "portafolios_por_n.csv")
    pesos = pd.read_csv(RES / "pesos_portafolios.csv")
    cov = pd.read_csv(RAIZ / "datos" / "yahoo" / "covarianza_1a.csv", index_col=0)
    g = res.loc[res.ganador].iloc[0]
    ganador = g.portafolio
    bal = "Balanceado (ganancia/riesgo) · 5 emisoras"
    ref = "Tu idea (NVDA + Banorte)"
    d = e.set_index("emisora")

    wb = load_workbook(ORIGINAL)
    CAP, COM, SES = "'Mi portafolio'!$C$4", "'Mi portafolio'!$C$5", "'Mi portafolio'!$H$4"

    # ---------------- Acciones: precios y estadísticas actualizadas + columnas nuevas ----------------
    ac = wb["Acciones"]
    ac["A1"].value = f"Las 145 acciones del Reto: mercado y datos históricos (precios actualizados: {HORA})"
    ac["A2"].value = ("Último precio, rendimientos, volatilidad, beta y ventanas: Yahoo Finance al momento indicado (SIC = precio en "
                      "dólares × tipo de cambio). Precio de compra/venta, diferencia y volumen: foto del monitor de la plataforma. "
                      "No es una recomendación de inversión.")
    act = {3: "ultimo", 4: "var_dia", 11: "r1m", 12: "r3m", 13: "r6m", 14: "r12m", 15: "vol_anual", 16: "mov_1s",
           17: "peor_caida", 18: "beta", 19: "mejor_vent", 20: "peor_vent", 21: "perfil"}
    nuevas = [("Precio objetivo analistas ($)", "objetivo_medio_mxn", "#,##0.00"), ("Potencial al objetivo", "potencial", PCT),
              ("Nº de analistas", "n_analistas", "0"), ("Próximo reporte", "prox_reporte", "General"),
              ("¿Reporta antes del 13-nov?", "reporta_en_reto", "General"),
              ("Importe operado diario ($, mexicanas)", "importe_diario_mxn", "#,##0"),
              ("Importe en mejor nivel del monitor ($)", "liq_monitor_mxn", "#,##0"), ("¿Líquida para el reto?", "liquida", "General"),
              ("Momento 12-1", "mom_12_1", PCT), ("Puntaje (z)", "puntaje", "0.00"), ("Beta vs canasta", "beta_canasta", "0.00"),
              ("Alfa (25 ses.)", None, "0.00%"), ("Costo ida y vuelta", None, "0.00%"), ("Pronóstico 25 ses. (neto)", None, "0.00%")]
    c0 = 22
    for j, (t, _, _) in enumerate(nuevas):
        c = ac.cell(4, c0 + j, t)
        c.font, c.fill, c.alignment = F_ENC, REL_ENC, CENTRO
        ac.column_dimensions[L(c0 + j)].width = 13
    col = {k: L(c0 + j) for j, (_, k, _) in enumerate(nuevas) if k}
    cA, cC, cP = L(c0 + 11), L(c0 + 12), L(c0 + 13)
    for r in range(5, 150):
        em = ac.cell(r, 1).value
        if em not in d.index:
            continue
        x = d.loc[em]
        for cc, k in act.items():
            ac.cell(r, cc).value = valor(x[k])
        for j, (_, k, fmt) in enumerate(nuevas):
            c = ac.cell(r, c0 + j)
            if k:
                c.value = valor(x[k])
            c.number_format, c.font, c.alignment = fmt, F_TXT, Alignment(horizontal="center")
        ac[f"{cA}{r}"] = f"='Pronóstico'!$B$5*P{r}*{col['puntaje']}{r}"
        ac[f"{cC}{r}"] = f"=IF(ISNUMBER(G{r}),G{r},1)+2*{COM}"
        ac[f"{cP}{r}"] = f"={col['beta_canasta']}{r}*'Pronóstico'!$B$8+{cA}{r}-{cC}{r}"
    ac.row_dimensions[4].height = 54
    rango = lambda k: f"Acciones!${k}$5:${k}$148"

    # ---------------- Covarianza y Correlaciones recalculadas ----------------
    cv, co = wb["Covarianza"], wb["Correlaciones"]
    nombres = [c.value for c in cv[4][1:145]]
    assert nombres == list(cov.index), "orden de emisoras distinto"
    sd = np.sqrt(np.diag(cov.values))
    corr = cov.values / np.outer(sd, sd)
    for i in range(144):
        for j in range(144):
            cv.cell(6 + i, 2 + j).value = float(cov.values[i, j])
            co.cell(5 + i, 2 + j).value = round(float(corr[i, j]), 2)

    # ---------------- Portafolios (pesos editables, fórmulas) ----------------
    po = wb.create_sheet("Portafolios")
    titulo(po, "Portafolios: ganador, alternativa balanceada y tu idea",
           "Pesos en celdas amarillas (se pueden cambiar). Pronóstico, alfa y costo salen de la hoja Acciones.")
    cart = [ganador, bal, ref]
    enc(po, 4, ["Emisora", "Pronóstico 25 ses.", "Beta vs canasta", "Alfa", "Costo ida y vuelta",
                "Ganador: " + ganador, "Alternativa: " + bal, ref])
    po.row_dimensions[4].height = 54
    union = list(dict.fromkeys(pesos[pesos.portafolio.isin(cart)].emisora))
    mapa = {(x.portafolio, x.emisora): x.peso for x in pesos.itertuples()}
    p0, p1 = 5, 4 + len(union)
    for i, em in enumerate(union):
        rr = p0 + i
        po.cell(rr, 1, em).font = F_TXT
        m = f"MATCH($A{rr},Acciones!$A$5:$A$148,0)"
        for cc, k, fmt in [(2, cP, "0.00%"), (3, col["beta_canasta"], "0.00"), (4, cA, "0.00%"), (5, cC, "0.00%")]:
            c = po.cell(rr, cc, f"=INDEX({rango(k)},{m})")
            c.number_format, c.font, c.alignment = fmt, F_TXT, Alignment(horizontal="center")
        for j, k in enumerate(cart):
            c = po.cell(rr, 6 + j, mapa.get((k, em), 0))
            c.number_format, c.font, c.fill, c.alignment = PCT, F_TXT, REL_IN, Alignment(horizontal="center")
    m0 = p1 + 2
    filas_m = ["Suma de pesos", "Emisoras distintas (mín. 5)", "Peso máximo (máx. 50%)", "Revisión de reglas",
               "Beta vs canasta", "Costo estimado ida y vuelta ($)", "Ganancia neta si el mercado baja ($)",
               "Ganancia neta, escenario base ($)", "Ganancia neta si el mercado sube ($)", "Ganancia esperada ($)"]
    for k, t in enumerate(filas_m):
        po.cell(m0 + k, 1, t).font = F_NEG
    for j in range(len(cart)):
        cc = L(6 + j)
        w = f"{cc}${p0}:{cc}${p1}"

        def esc(fila_r, cc=cc, w=w):
            return (f"={CAP}*(SUMPRODUCT({w},$C${p0}:$C${p1})*'Pronóstico'!$C${fila_r}"
                    f"+SUMPRODUCT({w},$D${p0}:$D${p1}))-{cc}{m0 + 5}")
        fm = [(f"=SUM({w})", PCT), (f'=COUNTIF({w},">0")', "0"), (f"=MAX({w})", PCT),
              (f'=IF(AND(ABS({cc}{m0}-1)<0.0005,{cc}{m0 + 1}>=5,{cc}{m0 + 2}<=0.5),"OK","REVISAR")', "General"),
              (f"=SUMPRODUCT({w},$C${p0}:$C${p1})", "0.00"), (f"=SUMPRODUCT({w},$E${p0}:$E${p1})*{CAP}", PESOS),
              (esc(13), PESOS), (esc(14), PESOS), (esc(15), PESOS),
              (f"='Pronóstico'!$B$13*{cc}{m0 + 6}+'Pronóstico'!$B$14*{cc}{m0 + 7}+'Pronóstico'!$B$15*{cc}{m0 + 8}", PESOS)]
        for k, (f, fmt) in enumerate(fm):
            put(po, f"{cc}{m0 + k}", f, fmt, F_NEG)
    po.column_dimensions["A"].width = 38
    for cc in "BCDE":
        po.column_dimensions[cc].width = 12
    for cc in "FGH":
        po.column_dimensions[cc].width = 26
    po.freeze_panes = "F5"

    # ---------------- Pronóstico ----------------
    pr = wb.create_sheet("Pronóstico")
    titulo(pr, "Pronóstico, portafolios de 5 a 12 emisoras y ganador",
           f"Datos al {HORA}. El pronóstico usa señales reales (precio objetivo de analistas y momento), pero su capacidad de predicción "
           "es baja: el resultado lo decide sobre todo el mercado. No es una recomendación de inversión.")
    pr["A4"].value, pr["A4"].font = "Supuestos (celdas amarillas editables)", F_SEC
    sup = [("IC: capacidad de pronóstico de las señales", 0.05, "0.00", "0.05 = una buena señal; 0 = sin pronóstico propio por acción"),
           ("Rendimiento anual esperado de la canasta", 0.08, "0.0%", "Prima de mercado"),
           ("Sesiones que faltan", f"={SES}", "0", "De la hoja Mi portafolio"),
           ("Rendimiento esperado de la canasta en el lapso", "=B6*B7/252", PCT, ""),
           ("Mov. típico (±1σ) de la canasta en el lapso", "=SQRT(SUM(Covarianza!$B$6:$EO$149)/144^2*B7/252)", "0.0%",
            "Canasta de las 144, mismo peso"),
           ("Umbral para el top 3", 0.20, "0%", "En una edición pasada los 3 primeros superaron +20%")]
    for k, (t, v, fmt, nota) in enumerate(sup):
        r = 5 + k
        put(pr, f"A{r}", t, None, F_NEG, centro=False)
        put(pr, f"B{r}", v, fmt, F_TXT, None if str(v).startswith("=") else REL_IN)
        put(pr, f"C{r}", nota, None, F_NOTA, centro=False)
    enc(pr, 12, ["Escenario", "Probabilidad", "Rend. de la canasta"])
    for k, (nom, p, rr) in enumerate([("Bajista", 0.25, "=B8-B9"), ("Base", 0.50, "=B8"), ("Alcista", 0.25, "=B8+B9")]):
        r = 13 + k
        put(pr, f"A{r}", nom, None, F_NEG, centro=False)
        put(pr, f"B{r}", p, "0%", F_TXT, REL_IN)
        put(pr, f"C{r}", rr, PCT)

    pr["A18"].value, pr["A18"].font = "¿Cuántas emisoras? Mejor portafolio para cada número (simulación de 60,000 trayectorias)*", F_SEC
    enc(pr, 19, ["Objetivo", "Emisoras", "Prob. de superar +20% (top 3)", "Prob. de superar +30%", "Ganancia esperada ($)",
                 "Prob. de terminar en positivo", "Prob. de perder más de $100k", "Peor 5% ($)", "Mejor 5% ($)", "Composición"])
    pr.row_dimensions[19].height = 42
    r = 20
    for x in res.itertuples():
        es_g = x.portafolio == ganador
        fill = REL_GAN if es_g else None
        put(pr, f"A{r}", x.objetivo, None, F_NEG if es_g else F_TXT, fill, centro=False)
        put(pr, f"B{r}", int(x.n), "0", F_FIJO, fill)
        for cc, v, fmt in [("C", x.P_top3_20pct, "0.0%"), ("D", x.P_30pct, "0.0%"), ("E", x.gan_esperada, PESOS),
                           ("F", x.P_ganar, "0.0%"), ("G", x.P_perder_100k, "0.0%"), ("H", x.P5, PESOS), ("I", x.P95, PESOS)]:
            put(pr, f"{cc}{r}", round(float(v), 4), fmt, F_FIJO, fill)
        put(pr, f"J{r}", x.emisoras, None, F_NOTA, centro=False)
        r += 1

    r += 1
    pr[f"A{r}"].value, pr[f"A{r}"].font = "Ganador", F_SEC
    rg = r + 1
    lineas = [("Portafolio ganador (máxima probabilidad de top 3)", ganador, None, F_NEG, REL_GAN),
              ("Composición", g.emisoras, None, F_TXT, None),
              ("Ganancia esperada con pronóstico ($)", f"='Portafolios'!F{m0 + 9}", PESOS, F_NEG, None),
              ("Si el mercado baja ($)", f"='Portafolios'!F{m0 + 6}", PESOS, F_TXT, None),
              ("Escenario base ($)", f"='Portafolios'!F{m0 + 7}", PESOS, F_TXT, None),
              ("Si el mercado sube ($)", f"='Portafolios'!F{m0 + 8}", PESOS, F_TXT, None),
              ("Ganancia esperada sin pronóstico propio, IC = 0 ($)*", round(float(g.gan_esperada_sin_alfa), 0), PESOS, F_FIJO, None),
              ("Prob. de superar +20% sin pronóstico propio*", round(float(g.P_top3_sin_alfa), 4), "0.0%", F_FIJO, None)]
    for k, (t, v, fmt, f, fill) in enumerate(lineas):
        put(pr, f"A{rg + k}", t, None, F_TXT, centro=False)
        put(pr, f"B{rg + k}", v, fmt, f, fill, centro=False)
    rn = rg + len(lineas) + 1
    notas = [
        "Cómo se pronostica cada acción: beta × rendimiento esperado de la canasta + alfa − costo. Alfa = IC × mov. típico × puntaje; el puntaje promedia el potencial al precio objetivo de analistas (5 o más analistas) y el momento de 12 meses sin el último mes, estandarizados entre las 144.",
        "Por qué gana un portafolio concentrado: el premio es para los 3 primeros, así que importa la probabilidad de un resultado muy alto, no el promedio. Más emisoras bajan el riesgo, pero también la probabilidad de llegar a +20%. El ganador mete el máximo permitido (50%) en la acción con mejor pronóstico y casi todo lo demás en la segunda; las otras 3 cumplen el mínimo de 5 emisoras.",
        "Universo: operables, diferencia compra-venta de 1.5% o menos y con liquidez (mexicanas: al menos $20 millones operados al día; SIC: al menos $40,000 en el mejor nivel del monitor). Para el objetivo de ganar solo entran emisoras con pronóstico positivo.",
        "* Valores grises (cursiva): simulación en Python con la covarianza del último año y colas gruesas (t de Student, 4 g.l.). No se recalculan si cambias los supuestos; las fórmulas sí.",
        "Riesgo: el ganador pierde más de $100,000 en cerca de 1 de cada 5 simulaciones. Es dinero virtual; los premios por avance (quizzes y Tracks) no dependen de esto.",
    ]
    for k, t in enumerate(notas):
        c = pr.cell(rn + k, 1, t)
        c.font, c.alignment = F_NOTA, Alignment(wrap_text=True, vertical="top")
        pr.merge_cells(start_row=rn + k, start_column=1, end_row=rn + k, end_column=10)
        pr.row_dimensions[rn + k].height = 30
    pr.column_dimensions["A"].width = 48
    pr.column_dimensions["B"].width = 16
    for j in range(3, 10):
        pr.column_dimensions[L(j)].width = 13
    pr.column_dimensions["J"].width = 70

    n_g = int((res.objetivo == g.objetivo).sum())
    ch = BarChart()
    ch.type, ch.grouping = "col", "clustered"
    ch.title = "Probabilidad de superar +20% según el número de emisoras"
    ch.add_data(Reference(pr, min_col=3, min_row=20, max_row=19 + n_g))
    ch.add_data(Reference(pr, min_col=3, min_row=20 + n_g, max_row=19 + 2 * n_g))
    ch.series[0].tx = SeriesLabel(v="Ganar (concentrado)")
    ch.series[1].tx = SeriesLabel(v="Balanceado")
    ch.set_categories(Reference(pr, min_col=2, min_row=20, max_row=19 + n_g))
    ch.y_axis.number_format = "0%"
    ch.y_axis.delete = False
    ch.x_axis.delete = False
    ch.x_axis.title = "Emisoras"
    ch.height, ch.width = 8, 18
    pr.add_chart(ch, f"A{rn + len(notas) + 1}")

    # ---------------- Órdenes del ganador ----------------
    od = wb.create_sheet("Órdenes")
    titulo(od, "Órdenes para el portafolio ganador",
           "Actualiza el precio (celda amarilla) con el precio de venta que veas en la plataforma justo antes de mandar la orden; "
           "los títulos se recalculan. Horario: 7:30 a 14:00 (desde el 3 de noviembre, 8:30 a 15:00). Las órdenes asignadas no se cancelan.")
    enc(od, 4, ["Emisora", "Tipo", "Peso", "Precio de venta ($)", "Títulos a comprar", "Importe ($)", "Comisión + IVA ($)",
                "Total ($)", "Próximo reporte"])
    od.row_dimensions[4].height = 36
    gw = pesos[pesos.portafolio == ganador].sort_values("peso", ascending=False)
    for k, x in enumerate(gw.itertuples()):
        rr = 5 + k
        m = f"MATCH($A{rr},Acciones!$A$5:$A$148,0)"
        put(od, f"A{rr}", x.emisora, None, F_NEG, centro=False)
        put(od, f"B{rr}", f"=INDEX(Acciones!$B$5:$B$148,{m})")
        put(od, f"C{rr}", float(x.peso), "0%", F_TXT, REL_IN)
        put(od, f"D{rr}", float(d.loc[x.emisora, "p_venta"]), "#,##0.00", F_TXT, REL_IN)
        put(od, f"E{rr}", f"=INT(C{rr}*{CAP}*0.997/D{rr})", "#,##0", F_NEG)
        put(od, f"F{rr}", f"=E{rr}*D{rr}", PESOS)
        put(od, f"G{rr}", f"=F{rr}*{COM}", "#,##0.00")
        put(od, f"H{rr}", f"=F{rr}+G{rr}", PESOS)
        put(od, f"I{rr}", f"=INDEX({rango(col['prox_reporte'])},{m})")
    u = 4 + len(gw)
    put(od, f"A{u + 1}", "Total", None, F_NEG, centro=False)
    for cc in "CFGH":
        put(od, f"{cc}{u + 1}", f"=SUM({cc}5:{cc}{u})", "0%" if cc == "C" else PESOS, F_NEG)
    put(od, f"A{u + 2}", "Efectivo que queda", None, F_NEG, centro=False)
    put(od, f"H{u + 2}", f"={CAP}-H{u + 1}", PESOS, F_NEG)
    pasos = ["Cómo operar:",
             "1. Revisa que el poder de compra sea $1,000,000 y que las 5 emisoras estén operables.",
             "2. Primero las dos posiciones grandes, con orden limitada al precio de venta que ves, para no pagar de más.",
             "3. Luego las tres chicas: cumplen el mínimo de 5 emisoras.",
             "4. Debe quedar un poco de efectivo para comisiones. No hace falta mover nada después: cada compra y venta cuesta.",
             "5. Si una orden no se asigna, vuelve a revisar el precio de venta: el monitor cambia durante el día."]
    for k, t in enumerate(pasos):
        od.cell(u + 4 + k, 1, t).font = F_NEG if k == 0 else F_TXT
    od.column_dimensions["A"].width = 16
    od.column_dimensions["B"].width = 14
    for cc in "CDEFGHI":
        od.column_dimensions[cc].width = 15

    # ---------------- Mi portafolio con el ganador ----------------
    mp = wb["Mi portafolio"]
    for k, x in enumerate(gw.itertuples()):
        mp.cell(8 + k, 1, x.emisora)
        mp.cell(8 + k, 2, float(x.peso))
    dv = DataValidation(type="list", formula1="Acciones!$A$5:$A$148", allow_blank=True, showErrorMessage=True,
                        errorTitle="Emisora", error="Elige una emisora de la lista (hoja Acciones).")
    dv.add("A8:A19")
    mp.add_data_validation(dv)

    # ---------------- Leeme ----------------
    lm = wb["Leeme"]
    lm["B3"].value = (f"Precios actualizados con Yahoo Finance: {HORA}. Precio de compra/venta y volumen del monitor: foto del 2026-10-08. "
                      "Datos descriptivos del pasado: NO son una recomendación de inversión y no predicen lo que va a pasar.")
    extra = [("Hojas agregadas (versión mejorada)", F_SEC),
             ("• Pronóstico: supuestos, mejor portafolio para cada número de emisoras (5 a 12), ganador y escenarios.", F_TXT),
             ("• Órdenes: títulos a comprar del ganador; actualiza el precio antes de mandar cada orden.", F_TXT),
             ("• Portafolios: pesos del ganador, la alternativa balanceada y tu idea original, con su ganancia por escenario.", F_TXT),
             ("• Acciones: columnas nuevas a la derecha (precio objetivo de analistas, próximo reporte, liquidez, pronóstico).", F_TXT),
             (f"• Mi portafolio viene llenado con el ganador ({ganador}). No es una recomendación de inversión.", F_TXT)]
    for k, (t, f) in enumerate(extra):
        c = lm.cell(30 + k, 2, t)
        c.font, c.alignment = f, Alignment(wrap_text=True, vertical="top")

    orden = ["Leeme", "Pronóstico", "Órdenes", "Mi portafolio", "Portafolios", "Acciones", "Correlaciones", "Covarianza"]
    wb._sheets = [wb[s] for s in orden]
    wb.active = 1
    for s in wb:
        s.sheet_view.tabSelected = s.title == "Pronóstico"
    wb.save(SALIDA)
    print("Guardado", SALIDA, "| ganador:", ganador)


if __name__ == "__main__":
    main()
