"""Reto Actinver 2026: descarga precios actualizados (Yahoo Finance) de las 145 emisoras del monitor.

Escribe en datos/yahoo/:
  precios_cierre.csv   cierres diarios ajustados de 3 años en moneda original (+ USDMXN=X y ^MXX)
  volumen.csv          volumen diario (títulos) del mercado de origen
  info_emisoras.csv    ticker, precio actual en pesos, precio objetivo de analistas, nº de analistas, próxima fecha de reporte
Los precios de compra/venta, spread y volumen del monitor de la plataforma no están en Yahoo: se conservan del Excel.
"""
import time
import warnings
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import pandas as pd
import yfinance as yf

warnings.filterwarnings("ignore")
RAIZ = Path(__file__).resolve().parent.parent
XLSX = RAIZ / "datos" / "Reto_Actinver_145_acciones.xlsx"
DEST = RAIZ / "datos" / "yahoo"

ESPECIALES = {"BRKB *": "BRK-B", "OXY1 *": "OXY", "CCL1 N": "CCL", "AA1 *": "AA"}


def ticker_yahoo(emisora, tipo):
    if emisora in ESPECIALES:
        return ESPECIALES[emisora]
    base = emisora.replace(" *", "").replace(" N", "") if tipo != "Mexicana" else emisora
    if tipo == "Mexicana":
        return base.replace(" *", "").replace(" ", "") + ".MX"
    return base.strip()


def emisoras():
    ws = openpyxl.load_workbook(XLSX, data_only=True)["Acciones"]
    filas = [(ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value) for r in range(5, 150)]
    df = pd.DataFrame(filas, columns=["emisora", "tipo", "ultimo_excel"])
    df["ticker"] = [ticker_yahoo(e, t) for e, t in zip(df.emisora, df.tipo)]
    return df


def main():
    DEST.mkdir(parents=True, exist_ok=True)
    df = emisoras()
    tickers = list(df.ticker) + ["USDMXN=X", "^MXX"]
    bruto = yf.download(tickers, period="3y", interval="1d", auto_adjust=True, progress=False, threads=True)
    px = bruto["Close"]
    px.to_csv(DEST / "precios_cierre.csv")
    bruto["Volume"].to_csv(DEST / "volumen.csv")
    fx = px["USDMXN=X"].dropna().iloc[-1]

    filas = []
    for r in df.itertuples():
        info = {}
        for intento in range(3):
            try:
                info = yf.Ticker(r.ticker).info or {}
                break
            except Exception:
                time.sleep(2 + 2 * intento)
        precio = info.get("regularMarketPrice") or info.get("currentPrice")
        moneda = info.get("currency")
        mult = fx if moneda == "USD" else 1.0
        ts = info.get("earningsTimestampStart") or info.get("earningsTimestamp")
        filas.append({
            "emisora": r.emisora, "ticker": r.ticker, "moneda": moneda,
            "precio_actual_mxn": precio * mult if precio else None,
            "ultimo_excel": r.ultimo_excel,
            "objetivo_medio_mxn": info.get("targetMeanPrice") * mult if info.get("targetMeanPrice") else None,
            "n_analistas": info.get("numberOfAnalystOpinions"),
            "recomendacion": info.get("recommendationMean"),
            "prox_reporte": datetime.fromtimestamp(ts, timezone.utc).date().isoformat() if ts else None,
            "hora_precio": datetime.fromtimestamp(info["regularMarketTime"], timezone.utc).isoformat()
            if info.get("regularMarketTime") else None,
        })
        time.sleep(0.25)
    out = pd.DataFrame(filas)
    out["usdmxn"] = fx
    out.to_csv(DEST / "info_emisoras.csv", index=False)
    out["dif_vs_excel"] = out.precio_actual_mxn / out.ultimo_excel - 1
    print(f"USDMXN {fx:.4f}; historial: {px.shape[0]} días, {px.notna().iloc[-260:].all().sum()} series completas en el último año")
    print("Sin precio actual:", list(out[out.precio_actual_mxn.isna()].emisora))
    print("Diferencia > 8% vs Excel (revisar ticker):")
    print(out[out.dif_vs_excel.abs() > 0.08][["emisora", "ticker", "precio_actual_mxn", "ultimo_excel"]].to_string(index=False))


if __name__ == "__main__":
    main()
